"""Create a blind human-review batch from the held-out annual-report test split."""
from __future__ import annotations

import csv
import json
import random
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SOURCE_DIR = ROOT / "答辩评测" / "annual_reports_weak_training_repaired_20261004"
OUT_DIR = ROOT / "答辩评测" / "semantic_review_batch_20261004"
SEED = 20261004
TARGETS = {"quote_本期": 50, "quote_上期": 50, "growth_set": 50}


def read_rows(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def fact_field(row: dict, key: str) -> str:
    value = row.get(key)
    if value not in (None, ""):
        return str(value)
    text = str(row.get("fact_text", ""))
    prefix = f"{key}="
    for part in text.split("；"):
        if part.startswith(prefix):
            return part[len(prefix):]
    return ""


def choose_claims(rows: list[dict]) -> list[tuple[str, list[dict]]]:
    groups: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        groups[row["claim_id"]].append(row)
    buckets: dict[str, list[tuple[str, list[dict]]]] = defaultdict(list)
    for claim_id, group in groups.items():
        buckets[group[0]["task_type"]].append((claim_id, group))
    rng = random.Random(SEED)
    selected: list[tuple[str, list[dict]]] = []
    for task_type, target in TARGETS.items():
        candidates = buckets[task_type]
        if len(candidates) < target:
            raise RuntimeError(f"not enough {task_type}: {len(candidates)} < {target}")
        # Spread the review over companies, then randomize within a stable seed.
        by_company: dict[str, list[tuple[str, list[dict]]]] = defaultdict(list)
        for item in candidates:
            by_company[item[1][0]["company"]].append(item)
        for items in by_company.values():
            rng.shuffle(items)
        companies = list(by_company)
        rng.shuffle(companies)
        picked: list[tuple[str, list[dict]]] = []
        while companies and len(picked) < target:
            next_companies = []
            for company in companies:
                items = by_company[company]
                if items:
                    picked.append(items.pop())
                    if len(picked) >= target:
                        break
                if items:
                    next_companies.append(company)
            companies = next_companies
        rng.shuffle(picked)
        selected.extend(picked)
    return selected


def write_outputs(selected: list[tuple[str, list[dict]]]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    claims = []
    candidates = []
    for index, (claim_id, group) in enumerate(selected, start=1):
        head = group[0]
        review_id = f"SR-{index:03d}"
        task_label = {"quote_本期": "本期单值", "quote_上期": "上期单值", "growth_set": "增长率双来源"}[head["task_type"]]
        claims.append({
            "review_id": review_id,
            "task_type": head["task_type"],
            "task_label": task_label,
            "claim_text": head["claim_text"],
            "company": head["company"],
            "source_claim_id": head.get("source_claim_id", ""),
            "source_file_hint": head.get("source_corpus", ""),
            "candidate_count": len(group),
            "expected_source_count": head["expected_k"],
            "manual_source_ids": "",
            "manual_abstain": "",
            "manual_note": "",
        })
        for position, row in enumerate(group, start=1):
            candidates.append({
                "review_id": review_id,
                "candidate_position": position,
                "fact_id": row["fact_id"],
                "fact_text": row["fact_text"],
                "metric": fact_field(row, "metric"),
                "period": fact_field(row, "period"),
                "value": row.get("fact_value", ""),
                "unit": fact_field(row, "unit"),
                "scope": fact_field(row, "scope"),
                "source_file_hint": row.get("source_corpus", ""),
                "manual_is_source": "",
                "manual_note": "",
            })
    claim_fields = list(claims[0])
    candidate_fields = list(candidates[0])
    for name, data, fields in (
        ("review_claims.csv", claims, claim_fields),
        ("review_candidates.csv", candidates, candidate_fields),
    ):
        with (OUT_DIR / name).open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            writer.writerows(data)
    with (OUT_DIR / "review_claims.jsonl").open("w", encoding="utf-8") as stream:
        for row in claims:
            stream.write(json.dumps(row, ensure_ascii=False) + "\n")
    manifest = {
        "batch": "semantic_review_batch_20261004",
        "seed": SEED,
        "source": str(SOURCE_DIR / "test.jsonl"),
        "source_split": "held-out test companies from the repaired weak-label corpus",
        "blind_review": True,
        "programmatic_labels_removed": True,
        "claims": len(claims),
        "candidate_rows": len(candidates),
        "task_counts": {key: sum(row["task_type"] == key for row in claims) for key in TARGETS},
        "companies": len({row["company"] for row in claims}),
        "files": ["review_claims.csv", "review_candidates.csv", "review_claims.jsonl", "审阅说明.md"],
    }
    (OUT_DIR / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    instructions = f'''# 语义审阅批次（{len(claims)} 条）

这是一批盲审数据，来自当前训练语料中未参与训练的 test 公司。程序原标签已从交付表中移除，不能把任何候选位置当成默认答案。

## 审阅方法

1. 在 `review_claims.csv` 读取一条论断，记下 `review_id`。
2. 在 `review_candidates.csv` 筛选同一个 `review_id`，逐项查看事实的主体、指标、期间、数值、单位和口径。
3. 将正确事实的 `fact_id` 填入 `manual_source_ids`：单值题填 1 个，增长率题填本期和上期 2 个，多个 ID 用 `;` 分隔。
4. 如果事实不足以证明论断，或主体/期间/单位/口径存在歧义，将 `manual_abstain` 填 `1`，并在 `manual_note` 写明原因；不要强行选择。
5. `review_candidates.csv` 的 `manual_is_source` 可填 `1`/`0`，用于记录逐候选判断；最终以 `review_claims.csv` 的来源集合为准。

## 覆盖范围

- 本期单值论断：{TARGETS['quote_本期']} 条
- 上期单值论断：{TARGETS['quote_上期']} 条
- 增长率双来源论断：{TARGETS['growth_set']} 条
- 公司数：{manifest['companies']} 家；候选事实：{len(candidates)} 条

## 标签用途

审阅完成后，这批数据可作为独立人工语义 gold：一部分用于阈值校准，一部分用于最终对比 BGE、BERT 和任务路由。不能把审阅结果回填到原训练集后再用同一条数据做最终测试。
'''
    (OUT_DIR / "审阅说明.md").write_text(instructions, encoding="utf-8")

    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "审阅论断"
        ws.append(claim_fields)
        for row in claims:
            ws.append([row[key] for key in claim_fields])
        ws2 = wb.create_sheet("候选事实")
        ws2.append(candidate_fields)
        for row in candidates:
            ws2.append([row[key] for key in candidate_fields])
        ws3 = wb.create_sheet("说明")
        for line in instructions.splitlines():
            ws3.append([line])
        for sheet in (ws, ws2):
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
                sheet.column_dimensions[column[0].column_letter].width = width
        wb.save(OUT_DIR / "semantic_review_batch_20261004.xlsx")
        manifest["files"].append("semantic_review_batch_20261004.xlsx")
        (OUT_DIR / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    except ImportError:
        manifest["xlsx"] = "not_generated: openpyxl unavailable"
        (OUT_DIR / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


def main() -> int:
    rows = read_rows(SOURCE_DIR / "test.jsonl")
    write_outputs(choose_claims(rows))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
