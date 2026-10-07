"""Reproducible same-source comparison for the PDF conversion backends.

The benchmark compares conversion quality and the additional evidence package
that distinguishes 知链 from a one-shot online PDF-to-DOC service. It never
uploads the source file; all backends run locally.
"""
from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path

import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from zhilian.pdf_office_export import export_pdf_to_office
from zhilian.pdf_revised_export import export_docx_to_pdf


def _pdf_stats(path: Path) -> dict:
    with pymupdf.open(path) as document:
        texts = [page.get_text("text", sort=True) or "" for page in document]
        return {"pages": len(document), "chars": sum(map(len, texts)),
                "nonempty_pages": sum(bool(text.strip()) for text in texts)}


def _docx_stats(path: Path) -> dict:
    document = Document(path)
    return {"paragraphs": len(document.paragraphs), "tables": len(document.tables),
            "sections": len(document.sections), "bytes": path.stat().st_size}


def run_backend(source: Path, out: Path, backend: str) -> dict:
    destination = out / backend
    destination.mkdir(parents=True, exist_ok=True)
    started = time.perf_counter()
    manifest = export_pdf_to_office(source, destination, subject=source.stem,
                                    docx_backend=backend)
    conversion_seconds = time.perf_counter() - started
    docx = Path(manifest["word"])
    exported_pdf = destination / f"{source.stem}_{backend}_word_export.pdf"
    started_export = time.perf_counter()
    pdf_export = export_docx_to_pdf(docx, exported_pdf)
    export_seconds = time.perf_counter() - started_export
    revised_stats = _pdf_stats(exported_pdf)
    source_stats = _pdf_stats(source)
    return {
        "backend": backend,
        "conversion_seconds": round(conversion_seconds, 3),
        "word_export_seconds": round(export_seconds, 3),
        "source_pdf": source_stats,
        "revised_pdf": revised_stats,
        "page_delta": revised_stats["pages"] - source_stats["pages"],
        "page_delta_percent": round((revised_stats["pages"] - source_stats["pages"]) /
                                     max(source_stats["pages"], 1) * 100, 3),
        "word": _docx_stats(docx),
        "facts": manifest.get("facts", 0),
        "claim_candidates": manifest.get("claim_candidate_count", 0),
        "text_anchor_count": manifest.get("text_anchor_count", 0),
        "artifacts": ["正文提取.docx", "事实表候选.xlsx", "文字锚点.json",
                      "论断候选.jsonl", "转换清单.json"],
        "revised_pdf_artifact": {**revised_stats, "path": str(exported_pdf),
                                  "export_backend": pdf_export.get("backend")},
    }


def _write_svg(path: Path, rows: list[dict]) -> None:
    width, height = 1000, 520
    labels = [row["backend"] for row in rows]
    values = [abs(row["page_delta_percent"]) for row in rows]
    max_value = max(values + [1])
    bars = []
    for index, (label, value) in enumerate(zip(labels, values)):
        x = 180 + index * 350
        bar_height = int(300 * value / max_value) if value else 4
        y = 380 - bar_height
        bars.append(f'<rect x="{x}" y="{y}" width="160" height="{bar_height}" rx="8" fill="#176852"/>')
        bars.append(f'<text x="{x + 80}" y="420" text-anchor="middle" font-size="22">{label}</text>')
        bars.append(f'<text x="{x + 80}" y="{y - 12}" text-anchor="middle" font-size="20">{value:.2f}%</text>')
    svg = ('<svg xmlns="http://www.w3.org/2000/svg" width="1000" height="520" viewBox="0 0 1000 520">'
           '<rect width="100%" height="100%" fill="#fbfaf7"/> '
           '<text x="500" y="48" text-anchor="middle" font-size="28" font-weight="700">PDF 重导出页数偏差</text>'
           '<text x="500" y="82" text-anchor="middle" font-size="17">同一份年报 · 本地 Word 导出 · 绝对偏差</text>'
           '<line x1="120" y1="380" x2="900" y2="380" stroke="#6b7280"/> '
           + ''.join(bars) +
           '<text x="500" y="480" text-anchor="middle" font-size="16">页数偏差越低，越适合作为修订版 PDF 的可编辑源</text></svg>')
    path.write_text(svg, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    rows = []
    for backend in ("native", "pdf2docx"):
        rows.append(run_backend(args.source.resolve(), args.out.resolve(), backend))
    capability_matrix = [
        {"capability": "格式转换", "ordinary_online": "通常返回 DOC/DOCX", "zhilian": "pdf2docx 优先，缺失时回退 native"},
        {"capability": "原件证据", "ordinary_online": "依赖用户自行留存", "zhilian": "原 PDF 只读、SHA-256、页码坐标锚点"},
        {"capability": "表格数据", "ordinary_online": "通常停留在 Word 表格", "zhilian": "候选 Excel、表格行、交叉校验状态"},
        {"capability": "论断与事实关联", "ordinary_online": "不属于格式转换范围", "zhilian": "论断 JSONL + 规则/本地模型/API 分层"},
        {"capability": "数值裁决", "ordinary_online": "不属于格式转换范围", "zhilian": "确定性引擎计算，用户确认后修复"},
        {"capability": "修订追踪", "ordinary_online": "通常只有下载文件", "zhilian": "变更报告、确认记录、独立修订版 PDF"},
        {"capability": "数据边界", "ordinary_online": "本次未做第三方上传测试", "zhilian": "本地转换可不上传原 PDF，远程仅发限定字段"},
    ]
    payload = {
        "source": str(args.source.resolve()),
        "local_only": True,
        "rows": rows,
        "online_service_comparison": {
            "scope": "capability comparison, not an upload benchmark",
            "ordinary_online_pdf_to_doc": ["format conversion"],
            "zhilian": ["format conversion", "PDF SHA-256", "page-coordinate anchors",
                        "claim candidates", "fact candidate Excel", "table cross-check",
                        "user confirmation", "change report", "independent revised PDF"],
            "privacy": "source remains local in this benchmark; no online upload",
            "matrix": capability_matrix,
        },
        "interpretation": [
            "pdf2docx improves editable-document layout on this born-digital sample, but does not replace provenance extraction.",
            "The product advantage is the auditable closed loop after conversion, not claiming ownership of the pdf2docx algorithm.",
        ],
    }
    (args.out / "pdf_backend_comparison.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    with (args.out / "pdf_online_capability_matrix.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        stream.write("能力,普通在线PDF转Word服务,知链当前交付\n")
        for row in capability_matrix:
            stream.write(",".join(row[key].replace(",", "，") for key in ("capability", "ordinary_online", "zhilian")) + "\n")
    _write_svg(args.out / "pdf_backend_page_delta.svg", rows)
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
