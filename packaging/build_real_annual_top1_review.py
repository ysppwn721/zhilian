"""Build a small human-Gold sheet for real annual-report Top-1 evaluation.

The current real run only exposes twelve admitted facts.  This script selects
claims whose wording can plausibly be checked against those facts, creates a
fixed candidate set, and writes a review workbook plus JSONL.  It deliberately
does not infer the human answer; ``gold_refs`` remains empty until reviewed.
"""

from __future__ import annotations

import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / "output" / "final_real_annual_verified_20261006" / "000543_2018_皖能电力"
OUT = ROOT / "output" / "final_real_annual_verified_20261006" / "real_annual_top1_review"
OUT.mkdir(parents=True, exist_ok=True)

CLAIM_IDS = [
    "48d94f30a20205a650",  # 13.67 亿元
    "15568a60a229b167f1",  # 38.19% growth
    "3f6a48a96ea0367f21",  # 13.67 亿元
    "8fd6e1d7a6b952ca64",  # exact current value
    "0842785a0032193bf0",  # revenue 9.90% growth
    "135c07b908539cf0d6",  # revenue in 万元
]

INSTRUCTIONS = (
    "每条只根据原文和候选事实复核。单值题填一个 fact_id；增长题填两个 fact_id，"
    "顺序建议按 prior,current。若现有事实无法证明，review_status 填 reject，gold_refs 留空。"
)


def load() -> tuple[dict[str, dict], dict[str, dict]]:
    workspace = json.loads((RUN / "baseline_workspace.json").read_text(encoding="utf-8"))
    facts = {x["id"]: x for x in workspace["facts"]}
    claims = {x["id"]: x for x in workspace["claims"]}
    return facts, claims


def build() -> list[dict]:
    facts, claims = load()
    rng = random.Random(20261006)
    all_facts = list(facts.values())
    rows = []
    for cid in CLAIM_IDS:
        claim = claims[cid]
        # A fixed, complete candidate universe prevents the reviewer from
        # silently changing the denominator between claims.
        candidates = list(all_facts)
        rng.shuffle(candidates)
        rows.append({
            "claim_id": cid,
            "kind": claim["kind"],
            "original": claim["original"],
            "source_location": claim["location"],
            "reported": claim.get("spec", {}).get("reported"),
            "candidate_fact_ids": [x["id"] for x in candidates],
            "gold_refs": [],
            "review_status": "pending",
            "review_note": "",
        })
    return rows


def write_jsonl(rows: list[dict]) -> Path:
    path = OUT / "real_annual_top1_review.jsonl"
    path.write_text("\n".join(json.dumps(x, ensure_ascii=False) for x in rows) + "\n", encoding="utf-8")
    return path


def write_workbook(rows: list[dict], facts: dict[str, dict]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "复核表"
    headers = [
        "claim_id", "类型", "原文论断", "原文位置", "论断数值",
        "候选1 fact_id", "候选1摘要", "候选2 fact_id", "候选2摘要",
        "候选3 fact_id", "候选3摘要", "候选4 fact_id", "候选4摘要",
        "候选5 fact_id", "候选5摘要", "候选6 fact_id", "候选6摘要",
        "候选7 fact_id", "候选7摘要", "候选8 fact_id", "候选8摘要",
        "候选9 fact_id", "候选9摘要", "候选10 fact_id", "候选10摘要",
        "候选11 fact_id", "候选11摘要", "候选12 fact_id", "候选12摘要",
        "gold_refs（填写ID，增长题填两个）", "review_status（confirmed/reject）", "review_note",
    ]
    ws.append(headers)
    for row in rows:
        values = [row["claim_id"], row["kind"], row["original"], row["source_location"], row["reported"]]
        for fid in row["candidate_fact_ids"]:
            f = facts[fid]
            summary = f"{f['metric']} | {f['period']} | {f['value']} {f['unit']} | {f['scope']}"
            values.extend([fid, summary])
        values.extend(["", "pending", ""])
        ws.append(values)
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    widths = {"A": 24, "B": 10, "C": 42, "D": 16, "E": 14, "AD": 28, "AE": 22, "AF": 24}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for col in range(6, 30):
        ws.column_dimensions[get_column_letter(col)].width = 20 if col % 2 else 36
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 34

    guide = wb.create_sheet("填写说明")
    guide.append(["项目", "说明"])
    guide.append(["目标", "建立真实年报小规模人工 Gold，用于计算 Top-1/Top-2/MRR/拒答率。"])
    guide.append(["复核规则", INSTRUCTIONS])
    guide.append(["confirmed", "证据足够，gold_refs 填正确 fact_id；quote 填 1 个，growth 填 prior 和 current 两个。"])
    guide.append(["reject", "现有 12 条事实无法证明该论断，gold_refs 留空；该条进入拒答评测，不算模型答对。"])
    guide.append(["口径", "优先检查指标、期间、单位、主体/统计口径和数值；万元与元可按 1 万换算。"])
    guide.append(["范围", "本表只覆盖现有事实源能形成候选的 6 条论断，不代表 113 条全文。"])
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 100
    for cell in guide[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    path = OUT / "real_annual_top1_review.xlsx"
    wb.save(path)
    return path


def main() -> None:
    facts, _ = load()
    rows = build()
    j = write_jsonl(rows)
    x = write_workbook(rows, facts)
    md = OUT / "复核说明.md"
    md.write_text(
        "# 真实年报 Top-1 人工复核\n\n"
        f"本表来自《皖能电力》2018 年年报，固定复核 {len(rows)} 条论断、每条 12 个候选事实。\n\n"
        f"{INSTRUCTIONS}\n\n"
        "复核完成后不要改候选 fact_id；只填写 gold_refs、review_status、review_note。\n",
        encoding="utf-8",
    )
    print(json.dumps({"xlsx": str(x), "jsonl": str(j), "rows": len(rows), "facts": len(facts)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
