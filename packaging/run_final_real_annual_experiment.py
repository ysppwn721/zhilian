"""Exercise the product on real annual-report text; never edit source PDFs or Gold.

Summary tables are a bounded experimental fact source, not an automatic approval.
Only explicitly denominated, cross-checked monetary rows are admitted. Simulated
updates are kept separate from observations of the original report.
"""
from __future__ import annotations

import argparse
from collections import Counter
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
import os
from pathlib import Path
import re
import sys
import time

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from zhilian import agent, engine, llm  # noqa: E402
from zhilian.office import HEADERS, read_document  # noqa: E402
from zhilian.store import Store  # noqa: E402


def dump(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


def clean(text):
    return re.sub(r'\s+', '', str(text or ''))


def numeric(text):
    raw = clean(text).replace(',', '').rstrip('%')
    if re.fullmatch(r'-?\d+(?:\.\d+)?', raw):
        return Decimal(raw)
    return None


def summary_facts(pdf, source):
    import pdfplumber
    records, rejected = [], []
    with pdfplumber.open(source) as document:
        for page_number, page in enumerate(document.pages[:20], 1):
            text = pdf[page_number - 1].get_text()
            if not re.search(r'主要会计数据|主要财务指标', text):
                continue
            for table_number, table in enumerate(page.find_tables(), 1):
                rows = table.extract()
                if not rows or not any('营业收入' in clean(row[0]) for row in rows if row):
                    continue
                current = prior = change = None
                pending = None
                prefix = ''

                def flush():
                    nonlocal pending
                    if pending:
                        records.append(pending)
                    pending = None

                for row_number, row in enumerate(rows, 1):
                    cells = [clean(cell) for cell in row]
                    years = {i: int(m.group(1)) for i, cell in enumerate(cells)
                             if (m := re.fullmatch(r'(20\d{2})年?(?:末)?', cell))}
                    if len(years) >= 2:
                        flush()
                        ordered = sorted(years, key=lambda i: -years[i])
                        current, prior = ordered[:2]
                        if years[current] - years[prior] != 1:
                            current = prior = None
                        change = next((i for i, cell in enumerate(cells) if '增减' in cell), None)
                        prefix = ''
                        continue
                    if current is None or prior is None or max(current, prior) >= len(cells):
                        continue
                    label = cells[0]
                    cur, pre = numeric(cells[current]), numeric(cells[prior])
                    if cur is not None and pre is not None:
                        flush()
                        pending = {'label': prefix + label, 'current': str(cur), 'prior': str(pre),
                                   'change': cells[change] if change is not None and change < len(cells) else '',
                                   'page': page_number, 'table': table_number, 'row': row_number,
                                   'bbox': list(table.bbox), 'raw_rows': [row]}
                        prefix = ''
                    elif pending:
                        if label.startswith(('归属于', '经营活动', '基本每股', '稀释每股', '加权平均', '总资产')):
                            flush()
                            prefix = label
                        else:
                            pending['label'] += label
                            pending['raw_rows'].append(row)
                            if change is not None and change < len(cells) and cells[change]:
                                pending['change'] = cells[change]
                    else:
                        prefix += label
                flush()
    admitted = []
    for record in records:
        label = record['label']
        unit_match = re.search(r'[（(](亿元|万元|元)[）)]', label)
        reported = numeric(record['change'])
        cur, pre = Decimal(record['current']), Decimal(record['prior'])
        reason = ('单位未明确或不是货币指标' if not unit_match else
                  '没有完整增长比例，未通过三值交叉检查' if reported is None or pre <= 0 else '')
        calculated = ((cur / pre - 1) * 100).quantize(Decimal('.01'), rounding=ROUND_HALF_UP) if pre > 0 else None
        if not reason and abs(calculated - reported) > Decimal('.02'):
            reason = '三值交叉检查不一致'
        record.update(calculated_change=str(calculated), rejection=reason)
        if reason:
            rejected.append(record)
            continue
        metric = label[:unit_match.start()]
        if not metric:
            rejected.append(dict(record, rejection='指标名为空'))
            continue
        record.update(metric=metric, unit=unit_match.group(1),
                      scope_basis='主要会计数据摘要；实验口径为公司整体，需用户复核')
        admitted.append(record)
    return admitted, rejected


def one_report(source, output):
    import pymupdf as fitz
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from openpyxl import Workbook

    started = time.perf_counter()
    folder = output / source.stem
    folder.mkdir(parents=True, exist_ok=True)
    original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    with fitz.open(source) as pdf:
        records, rejected = summary_facts(pdf, source)
        dump(folder / 'table_evidence.json', {'admitted_rows': records, 'rejected_rows': rejected})
        if not records:
            result = {'source': source.name, 'source_sha256': original_hash,
                      'status': 'no_admissible_summary_table', 'rejected_rows': len(rejected),
                      'source_pdf_unchanged': original_hash == hashlib.sha256(source.read_bytes()).hexdigest()}
            dump(folder / 'result.json', result)
            return result
        wb = Workbook()
        ws = wb.active
        ws.title = '事实表'
        ws.append(HEADERS + ['PDF页码', '表格坐标', '三值校验', '口径依据'])
        for i, record in enumerate(records):
            for period, key in [('本期', 'current'), ('上期', 'prior')]:
                ws.append([f'f{i}_{key}', '总计', record['metric'], period, float(record[key]),
                           record['unit'], '公司整体', record['page'], json.dumps(record['bbox']),
                           '通过；不是独立审计意见', record['scope_basis']])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for col in 'ABCDEFGHIJK':
            ws.column_dimensions[col].width = 22 if col not in 'CHI' else 45
        excel = folder / (source.stem + '_事实源.xlsx')
        wb.save(excel)
        doc = Document()
        normal = doc.styles['Normal']
        normal.font.name = '宋体'
        normal.font.size = Pt(10.5)
        normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
        doc.add_heading(source.stem + ' 正文实验副本', 0)
        doc.add_paragraph('公开年报文字层的前30页实验副本；不承诺排版保真。原始PDF只读，事实表仅覆盖主要会计数据货币指标。')
        anchors = []
        for page_number, page in enumerate(pdf, 1):
            if page_number > 30:
                break
            doc.add_heading(f'原始PDF第{page_number}页', 1)
            for block in page.get_text('blocks', sort=True):
                if block[6] != 0:
                    continue
                original = block[4]
                normalized = re.sub(r'(?<=[\u3400-\u9fff])\s+(?=[\u3400-\u9fff])', '', original)
                normalized = re.sub(r'\s+', ' ', normalized).strip()
                if not normalized:
                    continue
                index = len(doc.paragraphs)
                doc.add_paragraph(normalized)
                anchors.append({'paragraph': index, 'page': page_number, 'bbox': list(block[:4]),
                                'original_text': original, 'word_text': normalized})
        word = folder / (source.stem + '_年度报告正文.docx')
        doc.save(word)
        pages = len(pdf)
    dump(folder / 'text_anchors.json', {'pdf_sha256': original_hash, 'anchors': anchors})
    store = Store(output / 'workspaces')
    workspace = store.create(source.stem + ' 真实年报实验', [excel, word])
    raw = store.read(workspace['id'])
    raw['model_mode'] = 'local'
    store.write(raw)
    baseline = {'summary': workspace['summary'], 'checks': workspace['checks']}
    baseline['extraction'] = dict(Counter(c['kind'] for c in workspace['claims']))
    baseline['numeric_blocks'] = sum(bool(re.search(r'\d+(?:\.\d+)?\s*(?:亿元|万元|元|%)', b['text'])) for b in raw['blocks'])
    with llm.using_mode('local'):
        run_start = time.perf_counter()
        workspace = agent.run_agent(store, workspace['id'], workspace['revision'])
        baseline['agent_seconds'] = round(time.perf_counter() - run_start, 3)
    dump(folder / 'baseline_workspace.json', workspace)
    baseline['summary'] = workspace['summary']
    baseline['checks'] = workspace['checks']
    profile = workspace.get('agent', {}).get('document_profile')
    routing = workspace.get('agent', {}).get('model_routing')
    anchor_checks = []
    for claim in workspace['claims']:
        if claim['kind'] == 'chart':
            continue
        location = json.loads(claim['location'])
        match = next((a for a in anchors if location[:1] == ['p'] and a['paragraph'] == location[1]), None)
        anchor_checks.append({'claim_id': claim['id'], 'pdf_page': match['page'] if match else None,
                              'bbox': match['bbox'] if match else None,
                              'matches_original': bool(match and claim['original'] == match['word_text'][claim['start']:claim['end']])})
    dump(folder / 'claim_anchor_checks.json', anchor_checks)
    # Confirm only rule-linked originals that already agree numerically. This is
    # an explicit experiment decision, not a product automatic approval.
    approved = [c for c in workspace['claims'] if c['refs'] and
                engine.check(c, workspace['facts'])['status'] == 'consistent']
    if approved:
        workspace = store.confirm(workspace['id'], workspace['revision'],
                                  [{'claim_id': c['id'], 'refs': c['refs']} for c in approved])
    changed_id = None
    for fact in workspace['facts']:
        if fact['period'] == '本期' and any(fact['id'] in c['refs'] for c in approved):
            changed_id = fact['id']
            changed_value = (Decimal(str(fact['value'])) * Decimal('.99')).quantize(Decimal('.01'))
            break
    recovery = {'executed': False}
    if changed_id:
        workspace = store.change(workspace['id'], workspace['revision'], {changed_id: str(changed_value)})
        pre_repair = store.read(workspace['id'])
        before_files = {d['id']: (store.folder(workspace['id']) / pre_repair['generation'] / d['stored_name']).read_bytes()
                        for d in pre_repair['documents']}
        repair_ids = [c['id'] for c in workspace['claims'] if c['confirmed'] and
                      engine.check(c, workspace['facts'])['status'] == 'inconsistent']
        if repair_ids:
            workspace = store.repair(workspace['id'], workspace['revision'], repair_ids)
            out = store.export_local(workspace['id'])
            raw = store.read(workspace['id'])
            repaired_word = next(d for d in raw['documents'] if d['kind'] == 'docx')
            blocks, _, _ = read_document(store.folder(raw['id']) / raw['generation'] / repaired_word['stored_name'],
                                          repaired_word['id'], raw['facts'])
            original_raw = pre_repair
            expected_locations = {c['location'] for c in original_raw['claims'] if c['id'] in repair_ids}
            before_map = {b['location']: b['text'] for b in original_raw['blocks']}
            unaffected_ok = all(before_map[b['location']] == b['text'] for b in blocks if b['location'] not in expected_locations)
            reopened_ok = all(engine.check(c, raw['facts'])['status'] == 'consistent'
                              for c in raw['claims'] if c['confirmed'])
            restored = store.undo(raw['id'], raw['revision'])
            restored_raw = store.read(restored['id'])
            undo_ok = all((store.folder(restored['id']) / restored_raw['generation'] / d['stored_name']).read_bytes() == before_files[d['id']]
                          for d in restored_raw['documents'])
            recovery = {'executed': True, 'changed_fact': changed_id, 'simulated_new_value': str(changed_value),
                        'repaired_claims': len(repair_ids), 'reread_confirmed_consistent': reopened_ok,
                        'unaffected_blocks_unchanged': unaffected_ok, 'undo_bytes_equal': undo_ok,
                        'export_folder': out['folder']}
    result = {'source': source.name, 'source_sha256': original_hash, 'source_pages': pages,
              'scope': '前30页文字层；只接入主要会计数据摘要货币事实',
              'status': 'ok', 'facts': len(records) * 2, 'admitted_table_rows': len(records),
              'rejected_table_rows': len(rejected), 'baseline': baseline,
              'profile': profile,
              'routing': routing,
              'anchors_checked': len(anchor_checks),
              'anchors_passed': sum(a['matches_original'] for a in anchor_checks),
              'experiment_confirmation_count': len(approved), 'simulated_update': recovery,
              'source_pdf_unchanged': original_hash == hashlib.sha256(source.read_bytes()).hexdigest(),
              'seconds': round(time.perf_counter() - started, 3),
              'accuracy': None, 'accuracy_note': '未建立这份完整正文的人工Gold；不将数值一致率称为抽取准确率'}
    dump(folder / 'result.json', result)
    return result


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--output', type=Path, default=ROOT / 'output/final_real_annual_20261006')
    ap.add_argument('--sources', nargs='*', default=['000539_2018_粤电力Ａ.pdf', '000543_2018_皖能电力.pdf',
                                                   '000586_2018_修订版_汇源通信.pdf'])
    args = ap.parse_args()
    args.output = args.output.resolve()
    os.environ.setdefault('ZHILIAN_LOCAL_RERANKER_ENABLED', '1')
    os.environ.setdefault('ZHILIAN_ANNUAL_RERANKER_ENABLED', '1')
    os.environ['ZHILIAN_OUTPUT_DIR'] = str(args.output / 'exported')
    args.output.mkdir(parents=True, exist_ok=True)
    results = []
    for name in args.sources:
        source = ROOT / '答辩评测/annual_reports_v3_pool' / name
        print('RUN', source.name, flush=True)
        try:
            result = one_report(source, args.output)
        except Exception as exc:
            result = {'source': name, 'status': 'failed', 'error': str(exc)}
        results.append(result)
        print(json.dumps({k: v for k, v in result.items() if k in {'source', 'status', 'facts', 'seconds', 'simulated_update', 'error'}}, ensure_ascii=False), flush=True)
        dump(args.output / 'summary.json', {'results': results, 'gold_modified': False, 'remote_api_calls': 0})
    return int(any(r['status'] == 'failed' for r in results))


if __name__ == '__main__':
    raise SystemExit(main())
