"""Bounded OOXML adapters, stable locations, run-preserving patches and real charts."""
from __future__ import annotations

import hashlib
import json
import os
import shutil
import threading
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from docx import Document
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.parts.embeddedpackage import EmbeddedXlsxPart

from .engine import stable_id, number, fmt, convert

HEADERS = ['事实ID', '主体', '指标', '期间', '数值', '单位', '统计口径']

# 只自动处理能与一组事实一一对应的原生图表。多系列、组合图和
# 3D 柱图仍然转人工，避免把系列之间的口径混在一起。
SUPPORTED_SINGLE_SERIES_CHARTS = {
    XL_CHART_TYPE.COLUMN_CLUSTERED,
    XL_CHART_TYPE.BAR_CLUSTERED,
    XL_CHART_TYPE.PIE,
    XL_CHART_TYPE.PIE_EXPLODED,
    XL_CHART_TYPE.THREE_D_PIE,
    XL_CHART_TYPE.THREE_D_PIE_EXPLODED,
    XL_CHART_TYPE.DOUGHNUT,
    XL_CHART_TYPE.DOUGHNUT_EXPLODED,
    XL_CHART_TYPE.LINE,
    XL_CHART_TYPE.LINE_MARKERS,
}


def chart_kind(chart_type):
    """Return the user-facing kind for a supported native chart."""
    if chart_type in {
        XL_CHART_TYPE.PIE, XL_CHART_TYPE.PIE_EXPLODED,
        XL_CHART_TYPE.THREE_D_PIE, XL_CHART_TYPE.THREE_D_PIE_EXPLODED,
        XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.DOUGHNUT_EXPLODED,
    }:
        return '饼图' if chart_type not in {XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.DOUGHNUT_EXPLODED} else '圆环图'
    if chart_type in {XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS}:
        return '折线图'
    return '条形图' if chart_type == XL_CHART_TYPE.BAR_CLUSTERED else '柱状图'


def _pptx_has_external_chart_links(path):
    """Detect chart externalData relationships before reading chart values."""
    try:
        with ZipFile(path) as archive:
            names = archive.namelist()
            rel_ns = {'r': 'http://schemas.openxmlformats.org/package/2006/relationships'}
            for rel_name in names:
                if not (rel_name.startswith('ppt/charts/_rels/') and rel_name.endswith('.rels')):
                    continue
                root = ET.fromstring(archive.read(rel_name))
                for rel in root.findall('r:Relationship', rel_ns):
                    target = (rel.attrib.get('Target') or '').replace('\\', '/')
                    if 'embeddings/' not in target and 'externalLinks/' in target:
                        return True
            return False
    except (BadZipFile, OSError):
        return False


# 摘要缓存。原实现每次调用都把整份文件读进内存重算 SHA256，而 store.verify()
# 会在**每次读取类请求**上对全部文档校验一遍（上限 500 份），于是"打开项目"这类
# 只读操作的开销与项目体积成正比。摘要只由文件内容决定，用 (路径, 大小, mtime_ns)
# 做键即可命中；任何写入都会改变 mtime/大小，因此不存在读到陈旧摘要的风险。
_DIGEST_CACHE = {}
_DIGEST_CACHE_MAX = 2048
_DIGEST_LOCK = threading.Lock()


def digest(path):
    path = Path(path)
    try:
        stat = path.stat()
        key = (os.path.normcase(str(path.resolve())), stat.st_size, stat.st_mtime_ns)
    except OSError:
        # 文件不可 stat（缺失或权限不足）时退化为直读，让原有异常语义保持不变。
        return hashlib.sha256(path.read_bytes()).hexdigest()
    with _DIGEST_LOCK:
        cached = _DIGEST_CACHE.get(key)
    if cached is not None:
        return cached
    value = hashlib.sha256(path.read_bytes()).hexdigest()
    with _DIGEST_LOCK:
        if len(_DIGEST_CACHE) >= _DIGEST_CACHE_MAX:
            # 简单整体清空：条目等价、重建廉价，避免引入 LRU 依赖与额外簿记。
            _DIGEST_CACHE.clear()
        _DIGEST_CACHE[key] = value
    return value


def validate_office(path):
    if Path(path).suffix.lower() not in ('.xlsx', '.docx', '.pptx'):
        raise ValueError('仅支持 xlsx、docx、pptx 文件')
    try:
        with ZipFile(path) as z:
            infos = z.infolist()
            if len(infos) > 5000 or sum(i.file_size for i in infos) > 80 * 1024 * 1024:
                raise ValueError('文件展开后过大，请拆分为较小文件')
            if any('vbaproject' in i.filename.lower() for i in infos):
                raise ValueError('首版不支持包含宏的文件')
            expected = {'.xlsx': 'xl/workbook.xml', '.docx': 'word/document.xml', '.pptx': 'ppt/presentation.xml'}[Path(path).suffix.lower()]
            if expected not in z.namelist():
                raise ValueError('扩展名与 Office 文件内容不一致')
    except BadZipFile:
        raise ValueError('文件不是有效的 Office 开放 XML 文件')


def read_facts(path, file_id, allow_empty=False):
    validate_office(path)
    wb = load_workbook(path, data_only=False, read_only=True)
    facts, used = [], set()
    try:
        for sheet in wb.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = list(next(rows, []))
            if not all(h in header for h in HEADERS):
                continue
            if sheet.max_row > 2002 or sheet.max_column > 100:
                raise ValueError('首版每张事实表限制为2000行、100列')
            idx = {h: header.index(h) for h in HEADERS}
            # 列字母只在循环外算一次。原先每行都调用 sheet.cell()，而
            # read_only 模式下每次 cell() 都要重新解析工作表 XML，整体是
            # O(行数²)：实测 500 行 8.6s、1000 行 55s、1999 行 173s，用户会
            # 以为程序卡死。改为列字母拼坐标后 1999 行约 0.12s，坐标字符串完全一致。
            cell_column = get_column_letter(idx['数值'] + 1)
            for rno, row in enumerate(rows, 2):
                if not any(v is not None for v in row):
                    continue
                raw = {h: row[i] if i < len(row) else None for h, i in idx.items()}
                fid = str(raw['事实ID'] or '').strip()
                if not fid or fid in used:
                    raise ValueError(f'{sheet.title} 第{rno}行：事实ID为空或重复')
                used.add(fid)
                if len(fid) > 80:
                    raise ValueError('事实ID不能超过80字符')
                for field in ('主体', '指标', '期间', '单位', '统计口径'):
                    if raw[field] is None or not str(raw[field]).strip():
                        raise ValueError(f'{sheet.title} 第{rno}行缺少{field}')
                value = raw['数值']
                if isinstance(value, str) and value.startswith('='):
                    raise ValueError(f'{sheet.title} 第{rno}行使用公式。请先重算并提供事实数值，首版不信任公式缓存')
                if value is not None:
                    value = float(number(value))
                facts.append({'id': fid, 'file_id': file_id, 'subject': str(raw['主体']).strip(),
                              'metric': str(raw['指标']).strip(), 'period': str(raw['期间']).strip(),
                              'value': value, 'unit': str(raw['单位']).strip(), 'scope': str(raw['统计口径']).strip(),
                              'sheet': sheet.title, 'cell': cell_column + str(rno)})
    finally:
        wb.close()
    if not facts and not allow_empty:
        raise ValueError('没有找到事实表。首行需包含：' + '、'.join(HEADERS) + '。可下载模板或载入演示项目')
    if len(facts) > 2000:
        raise ValueError('首版单个项目最多2000条事实')
    return facts


def inspect_fact_source(path):
    """Return a conservative quality report for an Excel fact source.

    Parsing remains strict for malformed schemas (``read_facts`` keeps the
    existing exception contract), while this report records quality warnings
    that are useful after import: hidden sheets, merged headers, missing
    numeric values, formulas, and conflicting semantic definitions.  A
    ``blocked`` report never permits automatic document repair.
    """
    path = Path(path)
    validate_office(path)
    issues = []
    seen_ids = {}
    semantic_values = {}
    fact_rows = 0
    recognized_sheets = 0

    def issue(code, severity, message, sheet=None, cell=None):
        item = {'code': code, 'severity': severity, 'message': message}
        if sheet:
            item['sheet'] = sheet
        if cell:
            item['cell'] = cell
        issues.append(item)

    def merged_refs(sheet):
        # Read-only worksheets intentionally omit merged_cells.  The OOXML
        # worksheet path is available and lets us inspect merges without
        # loading a second full workbook into memory.
        worksheet_path = getattr(sheet, '_worksheet_path', None)
        if not worksheet_path:
            return []
        try:
            with ZipFile(path) as archive:
                root = ET.fromstring(archive.read(worksheet_path))
            return [node.attrib.get('ref', '') for node in root.iter()
                    if node.tag.rsplit('}', 1)[-1] == 'mergeCell' and node.attrib.get('ref')]
        except (KeyError, ET.ParseError, BadZipFile):
            return []

    wb = load_workbook(path, data_only=False, read_only=True)
    try:
        for sheet in wb.worksheets:
            refs = merged_refs(sheet)
            if sheet.sheet_state != 'visible':
                issue('hidden_sheet', 'warning', '工作表处于隐藏状态，事实是否应参与核验需要确认', sheet.title)
            if refs:
                issue('merged_cells', 'warning', f'工作表包含 {len(refs)} 个合并区域，表头或数据定位需人工确认', sheet.title)
            rows = sheet.iter_rows(values_only=True)
            header = list(next(rows, []))
            if not all(h in header for h in HEADERS):
                core_header_hits = sum(h in header for h in ('主体', '指标', '期间', '数值'))
                if '事实ID' in header or core_header_hits >= 3:
                    issue('missing_headers', 'blocking', '首行未包含完整事实表字段：' + '、'.join(HEADERS), sheet.title)
                continue
            recognized_sheets += 1
            idx = {h: header.index(h) for h in HEADERS}
            value_column = get_column_letter(idx['数值'] + 1)
            for rno, row in enumerate(rows, 2):
                if not any(v is not None and str(v).strip() for v in row):
                    continue
                fact_rows += 1
                raw = {h: row[i] if i < len(row) else None for h, i in idx.items()}
                fid = str(raw['事实ID'] or '').strip()
                cell = value_column + str(rno)
                if not fid:
                    issue('missing_id', 'blocking', '事实ID为空', sheet.title, cell)
                elif fid in seen_ids:
                    issue('duplicate_id', 'blocking', f'事实ID重复：{fid}（首次位于 {seen_ids[fid]}）', sheet.title, cell)
                else:
                    seen_ids[fid] = f'{sheet.title}!{cell}'
                for field in ('主体', '指标', '期间', '单位', '统计口径'):
                    if raw[field] is None or not str(raw[field]).strip():
                        issue('missing_field', 'blocking', f'缺少{field}', sheet.title, cell)
                value = raw['数值']
                if value is None or (isinstance(value, str) and not value.strip()):
                    issue('missing_value', 'blocking', '数值为空，不能自动修复', sheet.title, cell)
                elif isinstance(value, str) and value.startswith('='):
                    issue('formula_value', 'blocking', '数值为公式，需先重算并提供缓存结果', sheet.title, cell)
                else:
                    try:
                        numeric = float(number(value))
                    except (TypeError, ValueError, ArithmeticError):
                        issue('invalid_value', 'blocking', f'数值无法解析：{value}', sheet.title, cell)
                        numeric = None
                    if numeric is not None:
                        semantic = tuple(str(raw[k] or '').strip() for k in ('主体', '指标', '期间', '单位', '统计口径'))
                        previous = semantic_values.get(semantic)
                        if previous is not None and previous[0] != numeric:
                            issue('conflicting_definition', 'blocking',
                                  f'相同主体/指标/期间/单位/口径出现不同数值（{previous[0]} 与 {numeric}）',
                                  sheet.title, cell)
                        else:
                            semantic_values[semantic] = (numeric, f'{sheet.title}!{cell}')
    finally:
        wb.close()
    if not recognized_sheets:
        issue('no_fact_table', 'blocking', '没有找到包含完整事实表字段的工作表')
    elif not fact_rows:
        issue('empty_fact_table', 'blocking', '事实表没有可用数据行；需人工补录或整理后才能核验')
    blocking = sum(i['severity'] == 'blocking' for i in issues)
    warnings = sum(i['severity'] == 'warning' for i in issues)
    status = 'blocked' if blocking else ('review' if warnings else 'pass')
    return {
        'status': status,
        'auto_repair_allowed': not blocking,
        'facts_scanned': fact_rows,
        'issues': issues,
        'summary': {'blocking': blocking, 'warnings': warnings,
                    'hidden_sheets': sum(i['code'] == 'hidden_sheet' for i in issues),
                    'merged_cells': sum(i['code'] == 'merged_cells' for i in issues),
                    'duplicate_ids': sum(i['code'] == 'duplicate_id' for i in issues),
                    'missing_values': sum(i['code'] == 'missing_value' for i in issues),
                    'formula_values': sum(i['code'] == 'formula_value' for i in issues)},
    }


def docx_paragraphs(doc):
    for i, para in enumerate(doc.paragraphs):
        yield ['p', i], f'正文第{i+1}段', para
    for ti, table in enumerate(doc.tables):
        seen = set()
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                for pi, para in enumerate(cell.paragraphs):
                    yield ['t', ti, ri, ci, pi], f'表{ti+1} 第{ri+1}行第{ci+1}列', para


def read_document(path, file_id, facts, max_blocks=2000):
    validate_office(path)
    blocks, charts, warnings = [], [], []
    ext = Path(path).suffix.lower()
    if ext == '.docx':
        doc = Document(path)
        entries = docx_paragraphs(doc)
        for loc, label, para in entries:
            if not para.text.strip():
                continue
            if ''.join(r.text for r in para.runs) != para.text:
                warnings.append(label + '含超链接或复杂域，未自动处理')
                continue
            blocks.append({'file_id': file_id, 'location': json.dumps(loc), 'label': label, 'text': para.text})
        if doc.inline_shapes:
            warnings.append('Word内的图片、嵌入图表与图片文字未纳入自动检查')
        # 逐段检查页眉页脚：只查 paragraphs[0] 时，"第一段空、第二段有字"的
        # 页眉不会产生任何警告，用户会误以为页眉已被核验。页眉里的表格同样
        # 不在自动检查范围内，出现即提示。
        def frame_has_text(frame):
            return any(p.text.strip() for p in frame.paragraphs) or bool(frame.tables)

        if any(frame_has_text(s.header) or frame_has_text(s.footer) for s in doc.sections):
            warnings.append('页眉页脚未纳入检查')
    elif ext == '.pptx':
        prs = Presentation(path)
        external_chart_links = _pptx_has_external_chart_links(path)
        for si, slide in enumerate(prs.slides):
            for sh in slide.shapes:
                if sh.has_text_frame:
                    for pi, para in enumerate(sh.text_frame.paragraphs):
                        if not para.text.strip():
                            continue
                        label = f'第{si+1}页 · {sh.name} · 第{pi+1}段'
                        if ''.join(r.text for r in para.runs) != para.text:
                            warnings.append(label + '包含软换行，未自动处理')
                            continue
                        blocks.append({'file_id': file_id, 'location': json.dumps(['p', si, sh.shape_id, pi]), 'label': label, 'text': para.text})
                if sh.has_table:
                    for ri, row in enumerate(sh.table.rows):
                        for ci, cell in enumerate(row.cells):
                            for pi, para in enumerate(cell.text_frame.paragraphs):
                                if para.text.strip() and ''.join(r.text for r in para.runs) == para.text:
                                    blocks.append({'file_id': file_id, 'location': json.dumps(['t', si, sh.shape_id, ri, ci, pi]),
                                                   'label': f'第{si+1}页表格 第{ri+1}行第{ci+1}列', 'text': para.text})
                if sh.has_chart:
                    chart = sh.chart
                    if external_chart_links:
                        warnings.append(f'第{si+1}页图表可能引用外部工作簿，未自动处理，需人工复核')
                        continue
                    if chart.chart_type not in SUPPORTED_SINGLE_SERIES_CHARTS:
                        warnings.append(f'第{si+1}页图表类型未纳入安全自动处理范围，需人工复核')
                        continue
                    if len(chart.plots) != 1:
                        warnings.append(f'第{si+1}页图表包含多个绘图区或组合图，未自动处理，需人工复核')
                        continue
                    categories = [str(c.label) for c in chart.plots[0].categories]
                    series_list = list(chart.series)
                    if len(series_list) > 1:
                        warnings.append(f'第{si+1}页检测到{len(series_list)}个图表系列，已拆成系列级候选，需逐系列人工确认')
                    for series_index, series in enumerate(series_list):
                        refs = []
                        for cat in categories:
                            options = [f for f in facts if f['metric'] in series.name and f['period'] == cat]
                            if len(options) == 1:
                                refs.append(options[0]['id'])
                        unit = next((f['unit'] for f in facts if f['id'] in refs), '')
                        values = list(series.values)
                        loc = json.dumps(['chart', si, sh.shape_id, series_index])
                        valid = len(refs) == len(values) and all(v is not None for v in values)
                        if not valid:
                            refs = []
                            unit = ''
                        label = (f'第{si+1}页 · 原生{chart_kind(chart.chart_type)}'
                                 if len(series_list) == 1 else
                                 f'第{si+1}页 · 原生{chart_kind(chart.chart_type)} · 系列{series_index+1}')
                        rendered = []
                        for category, value in zip(categories, values):
                            rendered.append(f'{category} {fmt(value) if value is not None else "缺失"}{unit}')
                        charts.append({'id': stable_id(file_id, loc, 'chart'), 'file_id': file_id,
                                       'location': loc, 'label': label,
                                       'kind': 'chart', 'refs': refs,
                                       'original': series.name + '：' + '，'.join(rendered),
                                       'spec': {'series': series.name, 'series_index': series_index,
                                                'series_count': len(series_list), 'categories': categories,
                                                'values': values, 'unit': unit},
                                       'confirmed': False, 'extraction': '图表结构识别·系列级候选',
                                       'issue': '' if valid else '系列无法唯一对应事实，需人工指定完整来源',
                                       'start': 0, 'end': 0})
                if sh.shape_type in (6, 13):
                    warnings.append(f'第{si+1}页包含组合对象或图片，未检查其内部内容')
        warnings.append('幻灯片备注及母版中的文字未纳入检查')
    else:
        raise ValueError('成果文件只支持Word和PPT')
    if len(blocks) > max_blocks:
        raise ValueError(f'成果文件内容超过当前限制（{max_blocks}个文本块）')
    return blocks, charts, warnings


def read_images(path, file_id):
    """提取 Word 内嵌图片和 PPT 顶层图片，返回描述符及读取警告。"""
    validate_office(path)
    images, warnings = [], []

    def add(part, location, label):
        images.append({'id': 'img' + uuid4().hex[:16], 'file_id': file_id,
                       'location': json.dumps(location), 'label': label,
                       'mime': part.content_type, 'blob': part.blob})

    if Path(path).suffix.lower() == '.docx':
        doc = Document(path)
        for index, shape in enumerate(doc.inline_shapes, 1):
            try:
                rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
                add(doc.part.related_parts[rid], ['inline', index], f'正文 · 内嵌图片{index}')
            except (AttributeError, KeyError, ValueError, TypeError):
                warnings.append(f'第 {index} 个内嵌图片无法读取，已跳过')
    elif Path(path).suffix.lower() == '.pptx':
        for si, slide in enumerate(Presentation(path).slides):
            for shape in slide.shapes:
                if shape.shape_type == 13:
                    try:
                        add(shape.image, ['image', si, shape.shape_id], f'第{si+1}页 · 图片文字（{shape.name}）')
                    except (AttributeError, KeyError, ValueError, TypeError, OSError):
                        warnings.append(f'第{si+1}页图片无法读取，已跳过')
                elif shape.shape_type == 6:
                    warnings.append(f'第{si+1}页包含组合对象，嵌套图片暂不处理')
    return images, warnings


def patch_runs(para, edits):
    """Replace character spans in reverse order, retaining surrounding run formatting."""
    for start, end, expected, replacement in sorted(edits, reverse=True):
        full = ''.join(r.text for r in para.runs)
        if full[start:end] != expected:
            raise ValueError('原文锚点已变化，已停止写入，请重新分析')
        offset, inserted = 0, False
        for run in para.runs:
            text = run.text
            left, right = offset, offset + len(text)
            offset = right
            if right <= start or left >= end:
                continue
            a, b = max(0, start-left), min(len(text), end-left)
            run.text = text[:a] + (replacement if not inserted else '') + text[b:]
            inserted = True
        if not inserted:
            raise ValueError('无法定位原文字符范围')


def apply_document(source, destination, patches, facts):
    ext = Path(source).suffix.lower()
    document = Document(source) if ext == '.docx' else Presentation(source)
    groups = {}
    for patch in patches:
        location = patch['location']
        if ext == '.pptx':
            parsed = json.loads(location)
            if parsed[0] == 'chart':
                # Group all series belonging to one chart together.
                location = json.dumps(parsed[:3])
        groups.setdefault(location, []).append(patch)
    byid = {f['id']: f for f in facts}
    for location, group in groups.items():
        loc = json.loads(location)
        if ext == '.docx':
            if loc[0] == 'p':
                para = document.paragraphs[loc[1]]
            else:
                para = document.tables[loc[1]].cell(loc[2], loc[3]).paragraphs[loc[4]]
        else:
            slide = document.slides[loc[1]]
            shape = next((s for s in slide.shapes if s.shape_id == loc[2]), None)
            if shape is None:
                raise ValueError('找不到原幻灯片对象')
            if loc[0] == 'chart':
                # A multi-series chart is represented by one claim per series,
                # but is written back in a single replace_data operation so an
                # approved series never erases its untouched siblings.
                chart_patches = group
                patch = chart_patches[0]
                s = patch['spec']
                series_index = int(s.get('series_index', loc[3] if len(loc) > 3 else 0))
                current_series = list(shape.chart.series)
                if series_index >= len(current_series):
                    raise ValueError('图表系列索引已变化，已停止写入')
                chart_data = CategoryChartData()
                chart_data.categories = s['categories']
                changed_by_index = {}
                for item in chart_patches:
                    item_spec = item['spec']
                    changed_by_index[int(item_spec.get('series_index', 0))] = item
                for index, current in enumerate(current_series):
                    item = changed_by_index.get(index)
                    if item is None:
                        values = list(current.values)
                        name = current.name
                    else:
                        item_spec = item['spec']
                        if len(item['refs']) != len(item_spec['categories']):
                            raise ValueError('图表系列来源数量与类别数量不一致')
                        values = [float(convert(byid[i]['value'], byid[i]['unit'], item_spec['unit']))
                                  for i in item['refs']]
                        name = item_spec['series']
                    chart_data.add_series(name, values)
                workbook = shape.chart.part.chart_workbook
                try:
                    workbook.xlsx_part
                except KeyError:
                    # 某些成果仅保留图表缓存，嵌入工作簿关系已失效；按已批准事实重建。
                    workbook.xlsx_part = EmbeddedXlsxPart.new(chart_data.xlsx_blob, shape.chart.part.package)
                shape.chart.replace_data(chart_data)
                continue
            para = shape.text_frame.paragraphs[loc[3]] if loc[0] == 'p' else shape.table.cell(loc[3], loc[4]).text_frame.paragraphs[loc[5]]
        edits = [(p['start'], p['end'], p['original'], p['expected']) for p in group]
        # Overlapping edits are unsafe even if their replacement happens to look correct.
        ordered = sorted(edits)
        if any(a[1] > b[0] for a, b in zip(ordered, ordered[1:])):
            raise ValueError('检测到重叠修改，需人工复核')
        patch_runs(para, edits)
    document.save(destination)


def _sheet_paths(archive):
    """工作表名 → zip 内路径（xl/worksheets/sheetN.xml）。"""
    import xml.etree.ElementTree as ET

    ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
          'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
          'pr': 'http://schemas.openxmlformats.org/package/2006/relationships'}
    workbook = ET.fromstring(archive.read('xl/workbook.xml'))
    rels = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
    target = {rel.get('Id'): rel.get('Target') for rel in rels.findall('pr:Relationship', ns)}
    result = {}
    sheets_node = workbook.find('m:sheets', ns)
    for index, sheet in enumerate(sheets_node if sheets_node is not None else []):
        name = sheet.get('name')
        rel_id = sheet.get('{%s}id' % ns['r'])
        path = (target.get(rel_id) or '').lstrip('/')
        if not path:
            path = 'xl/worksheets/sheet%d.xml' % (index + 1)
        elif not path.startswith('xl/'):
            path = 'xl/' + path
        result[name] = path
    return result


def update_workbook(source, destination, facts, changes):
    """只替换目标单元格的数值，其余 zip 条目原样拷贝。

    为什么不用 openpyxl 的 load→save：**openpyxl 只建模它支持的部分，图表、条件格式、
    数据透视表、图片等未建模内容在保存时会静默丢失**。实测：一份带柱状图的事实表经
    load_workbook + save 后图表消失，而导出是"成功"的——这是最糟糕的一类缺陷：
    用户的数据被悄悄改坏了，却没有任何提示。

    这里改为以 xlsx（本质是 zip）为单位做**定点改写**：只替换 changes 里那几个
    单元格的 <v> 值，其它 entry 逐字节拷贝。图表引用的是单元格区域，因此数值被正确
    改写后，**图表会在打开时自动重绘**——不需要我们理解图表本身。

    工作表名/单元格都来自本项目自己解析并持久化的事实记录，因此对 xlsx 结构的要求
    仅限于标准 OOXML 布局；遇到无法解析的结构时抛错，绝不静默产出残缺文件。
    """
    import re
    import shutil
    import zipfile

    wanted = {}
    for fact in facts:
        if fact['id'] in changes:
            wanted.setdefault(fact['sheet'], {})[fact['cell']] = float(number(changes[fact['id']]))
    if not wanted:
        shutil.copyfile(source, destination)
        return

    with zipfile.ZipFile(source) as archive:
        names = set(archive.namelist())
        sheets = _sheet_paths(archive)
        missing = [name for name in wanted if sheets.get(name) not in names]
        if missing:
            raise ValueError('事实表结构无法解析，未修改文件：找不到工作表 %s' % '、'.join(missing))

        patched = {}
        for sheet_name, cells in wanted.items():
            path = sheets[sheet_name]
            patched[path] = (path, _patch_sheet(archive.read(path), cells))

        with zipfile.ZipFile(destination, 'w', zipfile.ZIP_DEFLATED) as out:
            for item in archive.infolist():
                # 保留每个条目的元数据；只替换被改写的工作表。
                data = patched[item.filename][1] if item.filename in patched else archive.read(item.filename)
                out.writestr(item, data)


def _patch_sheet(xml_bytes, cells):
    """把工作表 XML 里指定单元格的数值替换掉，返回新的 XML 字节。"""
    import re

    text = xml_bytes.decode('utf-8')
    for ref, value in cells.items():
        # 数值统一用最短往返表示，避免 125.0 这类尾零进入文件。
        rendered = repr(float(value))
        if rendered.endswith('.0'):
            rendered = rendered[:-2]
        # 1) 已有值：<c r="E2" ...><v>旧值</v>...
        pattern_value = re.compile(r'(<c\s[^>]*\br="%s"[^>]*>)(.*?)(</c>)' % re.escape(ref), re.S)
        # 2) 空单元格：<c r="E2" .../>
        pattern_empty = re.compile(r'<c\s[^>]*\br="%s"[^>]*/>' % re.escape(ref))

        def replace(match):
            head, body, tail = match.group(1), match.group(2), match.group(3)
            # 去掉原有的 t="s" 等类型标记：我们要写入的是数值。
            head = re.sub(r'\s+t="[^"]*"', '', head)
            if '<v>' in body:
                body = re.sub(r'<v>.*?</v>', '<v>%s</v>' % rendered, body, count=1, flags=re.S)
            elif '<is>' in body:                     # 内联字符串单元格
                body = '<v>%s</v>' % rendered
            else:
                body = '<v>%s</v>' % rendered + body
            return head + body + tail

        new_text, count = pattern_value.subn(replace, text, count=1)
        if count:
            text = new_text
            continue
        new_text, count = pattern_empty.subn('<c r="%s"><v>%s</v></c>' % (ref, rendered), text, count=1)
        if count:
            text = new_text
            continue
        # 单元格不存在：不猜位置插入（会破坏行列顺序），直接报错让人知道。
        raise ValueError('事实表中找不到单元格 %s，未修改文件' % ref)
    return text.encode('utf-8')
