from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation

from zhilian.demo import create_demo
from zhilian.engine import check, extract_claims
from zhilian.office import read_facts
from zhilian.store import Store


@pytest.fixture
def project(tmp_path):
    store = Store(tmp_path / 'data')
    w = store.create('测试项目', create_demo(tmp_path / 'files'), True)
    return store, w


def confirm_all(store, w):
    return store.confirm(w['id'], w['revision'], [{'claim_id': c['id'], 'refs': c['refs']} for c in w['claims']])


def test_delete_project_removes_workspace_and_all_versions(project, tmp_path):
    store, w = project
    folder = store.folder(w['id'])
    assert folder.is_dir()
    result = store.delete(w['id'])
    assert result == {'id': w['id'], 'name': '测试项目', 'deleted': True}
    assert not folder.exists()
    assert store.list() == []
    with pytest.raises(FileNotFoundError):
        store.read(w['id'])


def test_real_files_roundtrip_and_undo(project):
    store, w = project
    assert w['summary']['claims'] >= 11
    assert w['summary']['inconsistent'] == 0
    w = confirm_all(store, w)
    w = store.change(w['id'], w['revision'], {'sales_current': 90, 'product_a': 60, 'spending': 110})
    assert w['summary']['inconsistent'] == 11
    ids = [c['id'] for c in w['claims'] if check(c, w['facts'])['status'] == 'inconsistent']
    w = store.repair(w['id'], w['revision'], ids)
    assert w['summary']['inconsistent'] == 0
    assert w['last_repair']['verified']
    state = store.read(w['id'])
    root = store.folder(w['id']) / state['generation']
    word = next(d for d in state['documents'] if d['kind'] == 'docx')
    doc = Document(root / word['stored_name'])
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert '本期销售额为90万元' in text
    assert '较上期下降10%' in text
    assert 'B产品销量最高' in text
    assert '支出超过预算' in text
    assert '本段是人工撰写的背景说明，修复时应保持原样。' in text
    para = next(p for p in doc.paragraphs if '本期销售额为90' in p.text)
    assert para.runs[0].bold is True
    # Same-text replacements retain all unaffected runs; inserted text adopts its first run.
    deck = next(d for d in state['documents'] if d['kind'] == 'pptx')
    prs = Presentation(root / deck['stored_name'])
    chart = next(s.chart for slide in prs.slides for s in slide.shapes if s.has_chart)
    assert list(chart.series[0].values) == [100, 90]
    archive = store.archive(w['id'])
    with ZipFile(archive) as z:
        assert {'业务数据.xlsx', '分析报告.docx', '业务汇报.pptx', '核验报告.md', '知链核验记录.json'} <= set(z.namelist())
    w = store.undo(w['id'], w['revision'])
    assert w['summary']['inconsistent'] == 11
    w = store.undo(w['id'], w['revision'])
    assert w['summary']['inconsistent'] == 0
    assert next(f for f in w['facts'] if f['id'] == 'sales_current')['value'] == 125


def test_changed_dependency_can_remain_true(project):
    store, w = project
    w = confirm_all(store, w)
    w = store.change(w['id'], w['revision'], {'sales_current': 123})
    thresholds = [check(c, w['facts']) for c in w['claims'] if c['kind'] == 'threshold']
    assert all(c['status'] == 'consistent' for c in thresholds)
    assert any(check(c, w['facts'])['status'] == 'inconsistent' for c in w['claims'] if c['kind'] == 'growth')


def test_unconfirmed_and_stale_revision_rejected(project):
    store, w = project
    w = store.change(w['id'], w['revision'], {'sales_current': 90})
    with pytest.raises(ValueError, match='来源已确认'):
        store.repair(w['id'], w['revision'], [w['claims'][0]['id']])
    with pytest.raises(ValueError, match='其他窗口'):
        store.change(w['id'], 0, {'sales_current': 80})


def test_zero_base_and_scope_conflict(project):
    _, w = project
    growth = next(c for c in w['claims'] if c['kind'] == 'growth')
    facts = deepcopy(w['facts'])
    next(f for f in facts if f['id'] == 'sales_prev')['value'] = 0
    assert check(growth, facts)['status'] == 'unverifiable'
    facts = deepcopy(w['facts'])
    next(f for f in facts if f['id'] == 'sales_current')['scope'] = '不同口径'
    assert check(growth, facts)['status'] == 'unverifiable'


def test_ambiguous_source_and_missing_evidence(project):
    _, w = project
    facts = deepcopy(w['facts'])
    duplicate = deepcopy(next(f for f in facts if f['id'] == 'sales_current'))
    duplicate.update(id='other-sales', scope='其他地区')
    claims = extract_claims({'file_id': 'x', 'location': '1', 'label': 'p', 'text': '本期销售额为125万元。'}, facts + [duplicate])
    assert claims[0]['refs'] == []
    assert check(claims[0], facts)['status'] == 'unverifiable'
    assert extract_claims({'file_id': 'x', 'location': '1', 'label': 'p', 'text': '增长主要来自渠道优化。'}, facts) == []


def test_tied_ranking(project):
    _, w = project
    rank = next(c for c in w['claims'] if c['kind'] == 'ranking')
    facts = deepcopy(w['facts'])
    next(f for f in facts if f['id'] == 'product_b')['value'] = 80
    r = check(rank, facts)
    assert r['status'] == 'inconsistent'
    assert r['expected'] == 'A产品、B产品销量并列最高'


def test_formula_rejected(tmp_path):
    path = create_demo(tmp_path)[0]
    wb = load_workbook(path)
    wb.active['E2'] = '=SUM(10,90)'
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match='公式'):
        read_facts(path, 'f')


def test_disk_tamper_rejected(project):
    store, w = project
    state = store.read(w['id'])
    path = store.folder(w['id']) / state['generation'] / state['documents'][0]['stored_name']
    with path.open('ab') as f:
        f.write(b'changed')
    with pytest.raises(ValueError, match='源文件'):
        store.change(w['id'], w['revision'], {'sales_current': 90})


def test_api_workflow_and_boundary(tmp_path, monkeypatch):
    monkeypatch.delenv('ZHILIAN_ACCESS_PASSWORD', raising=False)
    from zhilian.app import create_app
    client = TestClient(create_app(tmp_path / 'api'))
    assert client.get('/').status_code == 200
    assert client.get('/api/health').json()['status'] == 'ok'
    w = client.post('/api/projects/demo').json()
    assert w['summary']['consistent'] >= 11
    response = client.post(f'/api/projects/{w["id"]}/facts', json={'revision':0,'values':{'sales_current':90}}, headers={'Origin':'https://evil.example'})
    assert response.status_code == 403
    assert client.get('/api/projects/not-valid').status_code == 400
    assert client.get('/api/projects/'+'a'*32).status_code == 404
    assert client.get(f'/api/projects/{w["id"]}/export').status_code == 200
    deleted = client.delete(f'/api/projects/{w["id"]}')
    assert deleted.status_code == 200 and deleted.json()['deleted'] is True
    assert client.get(f'/api/projects/{w["id"]}').status_code == 404
    monkeypatch.setenv('ZHILIAN_ACCESS_PASSWORD', 'local-test-password')
    assert client.get('/api/health').status_code == 401
    assert client.get('/api/health', auth=('zhilian', 'local-test-password')).status_code == 200


def test_upload_and_persistence(tmp_path, monkeypatch):
    monkeypatch.delenv('ZHILIAN_ACCESS_PASSWORD', raising=False)
    from zhilian.app import create_app
    paths = create_demo(tmp_path / 'input')
    client = TestClient(create_app(tmp_path / 'data'))
    response = client.post('/api/projects', data={'name':'上传项目'}, files=[('files',(p.name,p.read_bytes(),'application/octet-stream')) for p in paths[:2]])
    assert response.status_code == 200, response.text
    w = response.json()
    assert w['name'] == '上传项目'
    assert w['summary']['documents'] == 2
    second_client = TestClient(create_app(tmp_path / 'data'))
    assert second_client.get(f'/api/projects/{w["id"]}').json()['summary']['claims'] == 5


def test_append_document_batch_keeps_one_excel_and_scans_new_files(tmp_path):
    store = Store(tmp_path / 'data')
    initial = create_demo(tmp_path / 'initial')
    w = store.create('批量追加测试', initial[:2])
    extra = create_demo(tmp_path / 'extra')[1]
    extra_copy = tmp_path / 'extra' / '第二份报告.docx'
    doc = Document(extra)
    doc.add_paragraph('追加批次文档中的新内容')
    doc.save(extra_copy)
    before = store.read(w['id'])
    updated = store.append_documents(w['id'], before['revision'], [extra_copy])
    state = store.read(w['id'])
    assert len([d for d in state['documents'] if d['kind'] == 'xlsx']) == 1
    assert len(state['documents']) == 3
    assert updated['batch']['status'] == 'done'
    new_id = next(d['id'] for d in state['documents'] if d['name'] == extra_copy.name)
    assert any(c['file_id'] == new_id for c in state['claims'])
    assert store.batch(w['id'], updated['batch']['id'])['count'] == 1
    with pytest.raises(ValueError, match='同名文件'):
        store.append_documents(w['id'], updated['revision'], [extra_copy])


def test_append_documents_api_rejects_excel_and_supports_batch(tmp_path, monkeypatch):
    monkeypatch.delenv('ZHILIAN_ACCESS_PASSWORD', raising=False)
    from zhilian.app import create_app
    app = create_app(tmp_path / 'api')
    client = TestClient(app)
    paths = create_demo(tmp_path / 'input')
    response = client.post('/api/projects', data={'name': '批次 API'},
                           files=[('files', (p.name, p.read_bytes(), 'application/octet-stream')) for p in paths[:2]])
    assert response.status_code == 200, response.text
    w = response.json()
    extra = tmp_path / 'input' / '追加报告.docx'
    doc = Document(paths[1])
    doc.add_paragraph('API 追加批次的新内容')
    doc.save(extra)
    response = client.post(f'/api/projects/{w["id"]}/documents', data={'revision': w['revision']},
                           files=[('files', (extra.name, extra.read_bytes(), 'application/octet-stream'))])
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['batch']['count'] == 1
    assert result['summary']['documents'] == 3
    batch = client.get(f'/api/projects/{w["id"]}/batches/{result["batch"]["id"]}')
    assert batch.status_code == 200 and batch.json()['status'] == 'done'
    excel = paths[0]
    rejected = client.post(f'/api/projects/{w["id"]}/documents', data={'revision': result['revision']},
                           files=[('files', (excel.name, excel.read_bytes(), 'application/octet-stream'))])
    assert rejected.status_code == 400


def test_append_documents_api_tolerant_mode_keeps_good_files_and_reports_bad(tmp_path, monkeypatch):
    monkeypatch.delenv('ZHILIAN_ACCESS_PASSWORD', raising=False)
    from zhilian.app import create_app
    app = create_app(tmp_path / 'api')
    client = TestClient(app)
    paths = create_demo(tmp_path / 'input')
    created = client.post('/api/projects', data={'name': '容错批次'},
                          files=[('files', (p.name, p.read_bytes(), 'application/octet-stream')) for p in paths[:2]])
    assert created.status_code == 200, created.text
    w = created.json()
    good = tmp_path / 'input' / '追加正常.docx'
    doc = Document(paths[1])
    doc.add_paragraph('容错模式应保留的正常文件')
    doc.save(good)
    bad = tmp_path / 'input' / '追加损坏.docx'
    bad.write_bytes(b'not an Office document')
    response = client.post(
        f"/api/projects/{w['id']}/documents",
        data={'revision': w['revision'], 'continue_on_error': 'true'},
        files=[
            ('files', (good.name, good.read_bytes(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')),
            ('files', (bad.name, bad.read_bytes(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')),
        ],
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['batch']['status'] == 'partial'
    assert result['batch']['count'] == 1 and result['batch']['rejected'] == 1
    assert result['batch_results']['accepted'][0]['name'] == good.name
    assert result['batch_results']['rejected'][0]['name'] == bad.name
    assert '有效的 Office' in result['batch_results']['rejected'][0]['reason']
    assert result['summary']['documents'] == 3
    assert result['revision'] == w['revision'] + 1


def test_append_documents_api_tolerant_mode_rejects_duplicate_but_accepts_new_file(tmp_path, monkeypatch):
    monkeypatch.delenv('ZHILIAN_ACCESS_PASSWORD', raising=False)
    from zhilian.app import create_app
    app = create_app(tmp_path / 'api')
    client = TestClient(app)
    paths = create_demo(tmp_path / 'input')
    created = client.post('/api/projects', data={'name': '重复文件隔离'},
                          files=[('files', (p.name, p.read_bytes(), 'application/octet-stream')) for p in paths[:2]])
    w = created.json()
    duplicate = paths[1]
    good = tmp_path / 'input' / '第三份报告.docx'
    doc = Document(paths[1])
    doc.add_paragraph('不应被重复文件失败拖累')
    doc.save(good)
    response = client.post(
        f"/api/projects/{w['id']}/documents",
        data={'revision': w['revision'], 'continue_on_error': 'true'},
        files=[
            ('files', (duplicate.name, duplicate.read_bytes(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')),
            ('files', (good.name, good.read_bytes(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')),
        ],
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['batch']['status'] == 'partial'
    assert result['batch']['count'] == 1 and result['batch']['rejected'] == 1
    assert result['batch_results']['rejected'][0]['name'] == duplicate.name
    assert result['batch_results']['accepted'][0]['name'] == good.name
    assert result['summary']['documents'] == 3


def test_append_documents_api_tolerant_mode_all_bad_does_not_change_project(tmp_path, monkeypatch):
    monkeypatch.delenv('ZHILIAN_ACCESS_PASSWORD', raising=False)
    from zhilian.app import create_app
    app = create_app(tmp_path / 'api')
    client = TestClient(app)
    paths = create_demo(tmp_path / 'input')
    created = client.post('/api/projects', data={'name': '全失败批次'},
                          files=[('files', (p.name, p.read_bytes(), 'application/octet-stream')) for p in paths[:2]])
    w = created.json()
    bad = tmp_path / 'input' / '全损坏.docx'
    bad.write_bytes(b'not an Office document')
    response = client.post(
        f"/api/projects/{w['id']}/documents",
        data={'revision': w['revision'], 'continue_on_error': 'true'},
        files=[('files', (bad.name, bad.read_bytes(), 'application/octet-stream'))],
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['batch']['status'] == 'failed'
    assert result['batch']['count'] == 0 and result['batch']['rejected'] == 1
    assert result['revision'] == w['revision']
    assert result['summary']['documents'] == w['summary']['documents']


def test_append_document_batch_rolls_back_when_one_file_is_invalid(tmp_path):
    store = Store(tmp_path / 'data')
    paths = create_demo(tmp_path / 'files')
    w = store.create('追加回滚测试', paths[:2])
    invalid = tmp_path / 'files' / '损坏.docx'
    invalid.write_bytes(b'not an Office document')
    before = store.read(w['id'])
    with pytest.raises(ValueError, match='有效的 Office'):
        store.append_documents(w['id'], before['revision'], [invalid])
    after = store.read(w['id'])
    assert after['revision'] == before['revision']
    assert [d['name'] for d in after['documents']] == [d['name'] for d in before['documents']]


def test_external_excel_import_repair_and_undo(tmp_path, monkeypatch):
    monkeypatch.delenv('ZHILIAN_ACCESS_PASSWORD', raising=False)
    from zhilian.app import create_app
    app = create_app(tmp_path / 'data')
    client = TestClient(app)
    w = client.post('/api/projects/demo').json()
    w = confirm_all(app.state.store, w)
    path = create_demo(tmp_path / 'external')[0]
    wb = load_workbook(path)
    wb.active['E3'] = 90
    wb.save(path)
    wb.close()
    endpoint = f'/api/projects/{w["id"]}/source'
    response = client.post(endpoint, data={'revision': w['revision']}, files={'file': (path.name, path.read_bytes())})
    assert response.status_code == 200, response.text
    updated = response.json()
    assert updated['summary']['inconsistent'] > 0
    assert all(c['confirmed'] for c in updated['claims'])
    assert next(f for f in updated['facts'] if f['id'] == 'sales_current')['value'] == 90
    assert any('125' in b['text'] for b in updated['blocks'])
    assert client.post(endpoint, data={'revision': w['revision']}, files={'file': (path.name, path.read_bytes())}).status_code == 400
    ids = [c['id'] for c in updated['claims'] if check(c, updated['facts'])['status'] == 'inconsistent']
    repaired = app.state.store.repair(w['id'], updated['revision'], ids)
    assert repaired['summary']['inconsistent'] == 0
    assert any('下降10%' in b['text'] for b in repaired['blocks'])
    restored = app.state.store.undo(w['id'], repaired['revision'])
    restored = app.state.store.undo(w['id'], restored['revision'])
    assert next(f for f in restored['facts'] if f['id'] == 'sales_current')['value'] == 125


def test_external_excel_semantic_change_rejected(project, tmp_path):
    store, w = project
    path = create_demo(tmp_path / 'external')[0]
    wb = load_workbook(path)
    wb.active['G3'] = '其他口径'
    wb.active['E3'] = 90
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match='统计口径'):
        store.import_source(w['id'], w['revision'], path)
    assert store.read(w['id'])['revision'] == w['revision']
