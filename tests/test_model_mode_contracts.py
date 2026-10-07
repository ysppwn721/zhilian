"""Persisted UI choices, offline boundaries and API-only comparisons."""
import json

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from zhilian import agent, app as app_module, llm, ocr, quota, reranker
from zhilian.demo import create_demo
from zhilian.store import Store


@pytest.fixture(autouse=True)
def isolated(monkeypatch):
    monkeypatch.setattr(app_module, 'load_environment', lambda: None)
    monkeypatch.setenv('ZHILIAN_ACCESS_PASSWORD', '')
    monkeypatch.setenv('ZHILIAN_LLM_MODE', 'hybrid')
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'test-mode-secret')
    monkeypatch.setenv('ZHILIAN_QUOTA_PER_CLIENT', '60')
    monkeypatch.setenv('ZHILIAN_QUOTA_GLOBAL', '600')
    monkeypatch.setattr(reranker, 'status', lambda: {'enabled': False})
    def no_http(*args, **kwargs):
        pytest.fail('Unexpected external model call')
    monkeypatch.setattr(httpx, 'post', no_http)


def project(tmp_path, semantic=True, extra=0):
    paths = create_demo(tmp_path / 'files', semantic_demo=semantic)
    if extra:
        doc = Document(paths[1])
        for _ in range(extra):
            doc.add_paragraph('公司本期实现营收125万元。')
        doc.save(paths[1])
    store = Store(tmp_path / 'data')
    return store, store.create('模式测试', paths, True)


def response(suggestions):
    return httpx.Response(200, json={'choices': [{'message': {
        'content': json.dumps({'suggestions': suggestions})}}]})


def test_default_and_scoped_modes_restore(monkeypatch):
    monkeypatch.delenv('ZHILIAN_LLM_MODE')
    assert llm.mode() == 'hybrid'
    with llm.using_mode('rules'):
        assert not llm.remote_allowed()
        assert llm.mode({'model_mode': 'api'}) == 'api'
    assert llm.mode() == 'hybrid'


@pytest.mark.parametrize('selected,local_count,remote_count', [
    ('rules', 0, 0), ('local', 1, 0), ('hybrid', 1, 0), ('api', 0, 1), ('api_only', 0, 11),
])
def test_five_routes_use_only_expected_tools(tmp_path, monkeypatch, selected, local_count, remote_count):
    store, ws = project(tmp_path)
    ws = store.set_model_mode(ws['id'], ws['revision'], selected)
    local_calls, remote_calls = [], []
    monkeypatch.setattr(reranker, 'status', lambda: {'enabled': True})
    def local(_store, _ws, payload):
        local_calls.extend(payload['claim_ids'])
        return {'suggestions': [{'claim_id': cid, 'refs': ['sales_current'], 'source': 'local-reranker'}
                                for cid in payload['claim_ids']], 'stats': {'zero_scored': len(payload['claim_ids'])}}
    def remote(_store, _ws, payload):
        remote_calls.extend(payload['claim_ids'])
        return {'suggestions': []}
    monkeypatch.setitem(agent.TOOLS, 'suggest_links_local', local)
    monkeypatch.setitem(agent.TOOLS, 'suggest_links_llm', remote)
    result = agent.run_agent(store, ws['id'], ws['revision'])
    route = result['agent']['model_routing']
    assert result['agent']['phase'] == 'awaiting_decision'
    assert route['mode'] == selected
    assert len(local_calls) == local_count
    assert len(remote_calls) == remote_count
    assert route['api_batches'] == (1 if remote_count else 0)
    assert quota.snapshot()['global_used'] == (1 if remote_count else 0)
    assert not any(c['confirmed'] for c in result['claims'])


@pytest.mark.parametrize('selected,expect_api', [('local', False), ('hybrid', True)])
def test_missing_model_fallback(tmp_path, monkeypatch, selected, expect_api):
    store, ws = project(tmp_path)
    ws = store.set_model_mode(ws['id'], ws['revision'], selected)
    remote_calls = []
    monkeypatch.setitem(agent.TOOLS, 'suggest_links_llm', lambda s,w,p: remote_calls.extend(p['claim_ids']) or {'suggestions': []})
    result = agent.run_agent(store, ws['id'], ws['revision'])
    assert bool(remote_calls) == expect_api
    assert result['agent']['model_routing']['local_unavailable']
    assert any(a['event'] == 'Agent 本地模型降级' for a in result['audit'])


def test_local_batches_cover_more_than_forty(tmp_path, monkeypatch):
    store, ws = project(tmp_path, extra=80)
    ws = store.set_model_mode(ws['id'], ws['revision'], 'local')
    batches = []
    monkeypatch.setattr(reranker, 'status', lambda: {'enabled': True})
    def local(s,w,p):
        batches.append(p['claim_ids'])
        return {'suggestions': [], 'stats': {'zero_scored': len(p['claim_ids']), 'abstained': len(p['claim_ids'])}}
    monkeypatch.setitem(agent.TOOLS, 'suggest_links_local', local)
    result = agent.run_agent(store, ws['id'], ws['revision'])
    route = result['agent']['model_routing']
    assert [len(b) for b in batches] == [40, 40, 1]
    assert route['local_zero_recall_claims'] == 81
    assert quota.snapshot()['global_used'] == 0


def test_api_only_clears_rule_refs_keeps_chart_and_manual_gate(tmp_path, monkeypatch):
    store, ws = project(tmp_path)
    ws = store.set_model_mode(ws['id'], ws['revision'], 'api_only')
    original = store.read(ws['id'])
    requests = []
    valid = next(c for c in original['claims'] if c['kind'] == 'quote')
    def post(url, **kwargs):
        payload = json.loads(kwargs['json']['messages'][1]['content'])
        requests.append(payload)
        assert all(c['refs'] == [] for c in payload['claims'])
        assert all(c['kind'] != 'chart' for c in payload['claims'])
        return response([{'claim_id': valid['id'], 'refs': ['sales_current'], 'reason': 'ok'},
                         {'claim_id': valid['id'], 'refs': ['invented'], 'reason': 'bad'}])
    monkeypatch.setattr(httpx, 'post', post)
    result = agent.run_agent(store, ws['id'], ws['revision'])
    assert len(requests) == 1
    candidates = result['agent']['candidates']
    assert len(candidates[valid['id']]['options']) == 1
    assert 'DeepSeek' in candidates[valid['id']]['options'][0]['reason']
    for claim in result['claims']:
        if claim['kind'] == 'chart':
            assert candidates[claim['id']]['options']
            assert claim['refs']
        else:
            assert claim['refs'] == []
            if claim['id'] != valid['id']:
                assert candidates[claim['id']]['options'] == []
    assert not any(c['confirmed'] for c in result['claims'])
    assert not result.get('last_repair')
    assert result['documents'] == ws['documents']
    with pytest.raises(ValueError, match='来源事实'):
        agent.decide(store, result['id'], result['revision'], [{
            'kind': result['agent']['pending']['kind'],
            'items': [{'claim_id': valid['id'], 'refs': ['invented']}]}])


def test_api_only_batches_all_text_and_keeps_quota(tmp_path, monkeypatch):
    store, ws = project(tmp_path, extra=80)
    ws = store.set_model_mode(ws['id'], ws['revision'], 'api_only')
    batches = []
    def post(url, **kwargs):
        batches.append(json.loads(kwargs['json']['messages'][1]['content'])['claims'])
        return response([])
    monkeypatch.setattr(httpx, 'post', post)
    result = agent.run_agent(store, ws['id'], ws['revision'])
    assert [len(b) for b in batches] == [40, 40, 11]
    assert result['agent']['model_routing']['api_batches'] == 3
    assert quota.snapshot()['global_used'] == 3


@pytest.mark.parametrize('failure', ['missing_key', 'timeout', 'quota'])
def test_api_only_failure_does_not_silently_add_rules(tmp_path, monkeypatch, failure):
    store, ws = project(tmp_path)
    ws = store.set_model_mode(ws['id'], ws['revision'], 'api_only')
    calls = []
    if failure == 'missing_key':
        monkeypatch.setenv('DEEPSEEK_API_KEY', '')
    elif failure == 'quota':
        monkeypatch.setenv('ZHILIAN_QUOTA_GLOBAL', '0')
    else:
        def fail(*args, **kwargs):
            calls.append(1)
            raise httpx.ReadTimeout('test-mode-secret')
        monkeypatch.setattr(httpx, 'post', fail)
    result = agent.run_agent(store, ws['id'], ws['revision'])
    assert result['agent']['phase'] == 'awaiting_decision'
    assert len(calls) == (1 if failure == 'timeout' else 0)
    assert all(not i['options'] for i in result['agent']['candidates'].values() if i['kind'] != 'chart')
    assert 'test-mode-secret' not in json.dumps(result)


def test_switch_persists_per_project_and_invalidates_pending(tmp_path):
    app = app_module.create_app(tmp_path / 'api')
    client = TestClient(app)
    ws = client.post('/api/projects/demo').json()
    other = client.post('/api/projects/demo').json()
    base = f'/api/projects/{ws["id"]}'
    waiting = client.post(base+'/agent/run', json={'revision': ws['revision']}).json()
    selected = client.post(base+'/model-mode', json={'revision': waiting['revision'], 'mode': 'rules'})
    assert selected.status_code == 200
    changed = selected.json()
    assert changed['model_mode'] == 'rules'
    assert changed['agent']['pending'] is None
    assert changed['documents'] == waiting['documents']
    assert changed['claims'] == waiting['claims']
    assert changed['agent']['trace'] == waiting['agent']['trace']
    assert not changed['suggestions']
    assert Store(app.state.store.root).read(ws['id'])['model_mode'] == 'rules'
    assert client.get(f'/api/projects/{other["id"]}').json()['model_mode'] == 'hybrid'
    assert client.get('/api/health').json()['model_mode'] == 'hybrid'
    assert client.post(base+'/model-mode', json={'revision': waiting['revision'], 'mode':'api'}).status_code == 400
    assert client.post(base+'/model-mode', json={'revision': changed['revision'], 'mode':'invalid'}).status_code == 422
    repeat=client.post(base+'/model-mode', json={'revision': changed['revision'], 'mode':'rules'}).json()
    assert repeat['revision'] == changed['revision']
    resumed=client.post(base+'/agent/run', json={'revision': changed['revision']}).json()
    assert resumed['agent']['run_id'] != waiting['agent']['run_id']
    assert resumed['agent']['model_routing']['mode'] == 'rules'


@pytest.mark.parametrize('selected', ['rules','local'])
def test_offline_modes_block_all_remote_entrypoints(tmp_path, selected):
    app=app_module.create_app(tmp_path/'api')
    client=TestClient(app)
    ws=client.post('/api/projects/demo').json()
    base=f'/api/projects/{ws["id"]}'
    ws=client.post(base+'/model-mode', json={'revision':ws['revision'],'mode':selected}).json()
    assert not ws['ocr_enabled']
    assert client.post(base+'/suggest', json={'revision':ws['revision']}).status_code == 400
    result=client.post(base+'/diagnosis/explain', json={'revision':ws['revision']}).json()
    assert result['diagnosis_explanations'] == {}
    with llm.using_mode(selected):
        assert llm.explain_diagnosis([{}]) == {}
        with pytest.raises(ValueError):
            llm.suggest_links(ws['claims'],ws['facts'])
        with pytest.raises(ValueError,match='OCR 未启用'):
            ocr.ocr_image(tmp_path/'nonexistent.png','image/png')
    assert quota.snapshot()['global_used'] == 0


def test_environment_mode_change_does_not_resume_old_route(tmp_path, monkeypatch):
    store, ws=project(tmp_path, semantic=False)
    waiting=agent.run_agent(store,ws['id'],ws['revision'])
    monkeypatch.setenv('ZHILIAN_LLM_MODE','rules')
    assert agent.agent_status(store,ws['id'])['stale']
    restarted=agent.run_agent(store,ws['id'],waiting['revision'])
    assert restarted['agent']['run_id'] != waiting['agent']['run_id']
    assert restarted['agent']['model_routing']['mode'] == 'rules'
