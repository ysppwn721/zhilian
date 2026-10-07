from openpyxl import Workbook, load_workbook

from zhilian.demo import create_demo
from zhilian.office import HEADERS, inspect_fact_source
from zhilian.store import Store


def workbook_with_fact(path, value=125):
    wb = Workbook()
    ws = wb.active
    ws.title = '事实表'
    ws.append(HEADERS)
    ws.append(['sales_current', '总计', '销售额', '本期', value, '万元', '演示业务'])
    wb.save(path)
    wb.close()
    return path


def test_clean_fact_source_passes_quality_gate(tmp_path):
    report = inspect_fact_source(workbook_with_fact(tmp_path / 'clean.xlsx'))
    assert report['status'] == 'pass'
    assert report['auto_repair_allowed'] is True
    assert report['issues'] == []


def test_non_fact_conversion_sheets_do_not_block_pdf_fact_quality(tmp_path):
    path = workbook_with_fact(tmp_path / 'with_conversion_sheets.xlsx')
    wb = load_workbook(path)
    wb.create_sheet('转换摘要').append(['项目', '值'])
    wb.create_sheet('表格原文').append(['来源文件', '来源页', '表格编号', '行号', '列号', '原始内容'])
    wb.create_sheet('表格校验').append(['来源文件', '来源页', '表格编号', '指标', '计算变动'])
    wb.save(path)
    wb.close()

    report = inspect_fact_source(path)

    assert report['status'] == 'pass'
    assert report['issues'] == []


def test_hidden_sheet_and_merged_cells_are_review_warnings(tmp_path):
    path = workbook_with_fact(tmp_path / 'review.xlsx')
    wb = load_workbook(path)
    ws = wb['事实表']
    ws.merge_cells('A10:B10')
    hidden = wb.create_sheet('隐藏草稿')
    hidden.append(HEADERS)
    hidden.append(['draft', '总计', '销售额', '本期', 125, '万元', '草稿'])
    hidden.sheet_state = 'hidden'
    wb.save(path)
    wb.close()
    report = inspect_fact_source(path)
    assert report['status'] == 'review'
    assert report['auto_repair_allowed'] is True
    assert report['summary']['hidden_sheets'] == 1
    assert report['summary']['merged_cells'] == 1


def test_missing_value_blocks_automatic_repair(tmp_path):
    report = inspect_fact_source(workbook_with_fact(tmp_path / 'missing.xlsx', value=None))
    assert report['status'] == 'blocked'
    assert report['auto_repair_allowed'] is False
    assert report['summary']['missing_values'] == 1


def test_conflicting_semantic_definition_blocks_automatic_repair(tmp_path):
    path = workbook_with_fact(tmp_path / 'conflict.xlsx')
    wb = load_workbook(path)
    wb['事实表'].append(['sales_current_copy', '总计', '销售额', '本期', 130, '万元', '演示业务'])
    wb.save(path)
    wb.close()
    report = inspect_fact_source(path)
    assert report['status'] == 'blocked'
    assert report['summary']['blocking'] >= 1
    assert any(i['code'] == 'conflicting_definition' for i in report['issues'])


def test_repair_is_refused_when_fact_quality_is_blocked(tmp_path):
    store = Store(tmp_path / 'data')
    w = store.create('质量门禁', create_demo(tmp_path / 'files'), True)
    w = store.confirm(w['id'], w['revision'], [{'claim_id': c['id'], 'refs': c['refs']} for c in w['claims']])
    w = store.change(w['id'], w['revision'], {'sales_current': 90})
    state = store.read(w['id'])
    state['fact_quality'] = {'status': 'blocked', 'auto_repair_allowed': False, 'issues': [{'code': 'missing_value'}]}
    store.write(state)
    inconsistent = [c['id'] for c in w['claims'] if c['confirmed']]
    try:
        store.repair(w['id'], w['revision'], inconsistent[:1])
    except ValueError as exc:
        assert '事实源质量' in str(exc)
    else:
        raise AssertionError('blocked fact quality must refuse automatic repair')
