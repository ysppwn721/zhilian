from pptx import Presentation
from pptx.chart.data import ChartData, CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from zhilian import office


def _facts():
    return [
        {'id': 'north', 'subject': '区域', 'metric': '销售额', 'period': '华北',
         'value': 60, 'unit': '万元', 'scope': '合计'},
        {'id': 'south', 'subject': '区域', 'metric': '销售额', 'period': '华南',
         'value': 40, 'unit': '万元', 'scope': '合计'},
    ]


def _make_deck(path, chart_type):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ['华北', '华南']
    data.add_series('销售额', [60, 40])
    slide.shapes.add_chart(chart_type, Inches(1), Inches(1), Inches(8), Inches(4), data)
    prs.save(path)


def _make_multi_series_deck(path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ['华北', '华南']
    data.add_series('销售额', [60, 40])
    data.add_series('成本额', [30, 20])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(8), Inches(4), data)
    prs.save(path)


def test_native_pie_chart_is_read_as_a_traceable_claim(tmp_path):
    path = tmp_path / 'pie.pptx'
    _make_deck(path, XL_CHART_TYPE.PIE)
    blocks, charts, warnings = office.read_document(path, 'P', _facts())
    assert not blocks
    assert not any('图表' in warning for warning in warnings)
    assert len(charts) == 1
    assert charts[0]['label'].endswith('原生饼图')
    assert charts[0]['refs'] == ['north', 'south']
    assert charts[0]['spec']['categories'] == ['华北', '华南']


def test_native_bar_chart_is_read_without_reclassifying_it_as_text(tmp_path):
    path = tmp_path / 'bar.pptx'
    _make_deck(path, XL_CHART_TYPE.BAR_CLUSTERED)
    _, charts, warnings = office.read_document(path, 'P', _facts())
    assert not any('图表' in warning for warning in warnings)
    assert charts[0]['label'].endswith('原生条形图')
    assert charts[0]['refs'] == ['north', 'south']


def test_native_pie_chart_data_is_rebuilt_after_fact_change(tmp_path):
    source = tmp_path / 'pie-source.pptx'
    output = tmp_path / 'pie-output.pptx'
    _make_deck(source, XL_CHART_TYPE.PIE)
    _, charts, _ = office.read_document(source, 'P', _facts())
    changed = [dict(f) for f in _facts()]
    changed[0]['value'] = 50
    patch = {'location': charts[0]['location'], 'spec': charts[0]['spec'], 'refs': charts[0]['refs']}
    office.apply_document(source, output, [patch], changed)
    prs = Presentation(output)
    chart = next(shape.chart for slide in prs.slides for shape in slide.shapes if shape.has_chart)
    assert list(chart.series[0].values) == [50, 40]


def test_multi_series_chart_is_split_into_series_level_candidates(tmp_path):
    path = tmp_path / 'multi.pptx'
    _make_multi_series_deck(path)
    facts = _facts() + [
        {'id': 'cost_north', 'subject': '区域', 'metric': '成本额', 'period': '华北',
         'value': 30, 'unit': '万元', 'scope': '合计'},
        {'id': 'cost_south', 'subject': '区域', 'metric': '成本额', 'period': '华南',
         'value': 20, 'unit': '万元', 'scope': '合计'},
    ]
    _, charts, warnings = office.read_document(path, 'P', facts)
    assert len(charts) == 2
    assert {c['spec']['series'] for c in charts} == {'销售额', '成本额'}
    assert all(len(c['refs']) == 2 for c in charts)
    assert any('系列级候选' in warning for warning in warnings)


def test_multi_series_repair_preserves_unselected_series(tmp_path):
    source = tmp_path / 'multi-source.pptx'
    output = tmp_path / 'multi-output.pptx'
    _make_multi_series_deck(source)
    facts = _facts() + [
        {'id': 'cost_north', 'subject': '区域', 'metric': '成本额', 'period': '华北',
         'value': 30, 'unit': '万元', 'scope': '合计'},
        {'id': 'cost_south', 'subject': '区域', 'metric': '成本额', 'period': '华南',
         'value': 20, 'unit': '万元', 'scope': '合计'},
    ]
    _, charts, _ = office.read_document(source, 'P', facts)
    sales = next(c for c in charts if c['spec']['series'] == '销售额')
    changed = [dict(f) for f in facts]
    next(f for f in changed if f['id'] == 'north')['value'] = 55
    patch = {'location': sales['location'], 'spec': sales['spec'], 'refs': sales['refs']}
    office.apply_document(source, output, [patch], changed)
    prs = Presentation(output)
    chart = next(shape.chart for slide in prs.slides for shape in slide.shapes if shape.has_chart)
    assert [list(s.values) for s in chart.series] == [[55, 40], [30, 20]]
