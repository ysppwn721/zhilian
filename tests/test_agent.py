import json
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation

from zhilian import agent, app as app_module
from zhilian.demo import create_demo
from zhilian.store import Store


@pytest.fixture(autouse=True)
def isolated_model(monkeypatch):
    # 禁止测试读取真实密钥或意外产生外部调用。
    monkeypatch.setattr(app_module, 'load_environment', lambda: None)
    monkeypatch.setenv('DEEPSEEK_API_KEY', '')
    monkeypatch.setenv('DEEPSEEK_MODEL', 'deepseek-flash')
    monkeypatch.setenv('ZHILIAN_ACCESS_PASSWORD', '')
    def unexpected(*args, **kwargs):
        pytest.fail('测试不应发起真实模型请求')
    monkeypatch.setattr(httpx, 'post', unexpected)


@pytest.fixture
def project(tmp_path):
    store = Store(tmp_path / 'data')
    return store, store.create('智能体测试', create_demo(tmp_path / 'files'), True)


def approve(store, ws, items=None):
    pending = ws['agent']['pending']
    return agent.decide(store, ws['id'], ws['revision'], [{
        'kind': pending['kind'], 'items': items if items is not None else [
            {'claim_id': i['claim_id'], 'refs': i['refs']} for i in pending['items']]}])


def file_bytes(store, wid):
    ws = store.read(wid)
    root = store.folder(wid) / ws['generation']
    return {d['id']: (root / d['stored_name']).read_bytes() for d in ws['documents']}


def test_offline_full_flow_modifies_real_office_and_undo(project):
    store, ws = project
    ws = store.change(ws['id'], ws['revision'], {'sales_current': 90, 'product_a': 60, 'spending': 110})
    before = file_bytes(store, ws['id'])
    ws = agent.run_agent(store, ws['id'], ws['revision'])
    assert ws['agent']['pending']['kind'] == 'confirm_links'
    assert not any(c['confirmed'] for c in ws['claims'])
    assert file_bytes(store, ws['id']) == before
    ws = approve(store, ws)
    assert ws['agent']['pending']['kind'] == 'approve_repair'
    assert all(c['confirmed'] for c in ws['claims'])
    assert file_bytes(store, ws['id']) == before
    ws = approve(store, ws)
    assert ws['agent']['phase'] == 'done'
    assert ws['agent']['result'] == {'consistent': 11, 'inconsistent': 0, 'unverifiable': 0, 'repaired': 11}
    raw = store.read(ws['id'])
    root = store.folder(ws['id']) / raw['generation']
    doc = Document(root / next(d['stored_name'] for d in raw['documents'] if d['kind'] == 'docx'))
    text = '\n'.join(p.text for p in doc.paragraphs)
    for expected in ('本期销售额为90万元', '较上期下降10%', 'B产品销量最高', '支出超过预算',
                     '本段是人工撰写的背景说明，修复时应保持原样。'):
        assert expected in text
    assert next(p for p in doc.paragraphs if '本期销售额为90' in p.text).runs[0].bold
    prs = Presentation(root / next(d['stored_name'] for d in raw['documents'] if d['kind'] == 'pptx'))
    chart = next(s.chart for slide in prs.slides for s in slide.shapes if s.has_chart)
    assert list(chart.series[0].values) == [100, 90]
    assert not any(t['model'] for t in ws['agent']['trace'])
    restored = store.undo(ws['id'], ws['revision'])
    assert restored['summary']['inconsistent'] == 11
    assert file_bytes(store, ws['id']) == before


def test_ambiguity_requires_explicit_choice(tmp_path):
    paths = create_demo(tmp_path / 'files')
    wb = load_workbook(paths[0])
    wb.active.append(['other_sales', '总计', '销售额', '本期', 125, '万元', '其他地区'])
    wb.save(paths[0])
    wb.close()
    store = Store(tmp_path / 'data')
    ws = store.create('歧义测试', paths)
    ws = agent.run_agent(store, ws['id'], ws['revision'])
    assert ws['agent']['pending']['kind'] == 'resolve_ambiguity'
    item = next(i for i in ws['agent']['pending']['items'] if i['kind'] == 'quote')
    assert len(item['options']) >= 2
    assert not next(c for c in ws['claims'] if c['id'] == item['claim_id'])['confirmed']
    ws = approve(store, ws, [{'claim_id': item['claim_id'], 'refs': ['sales_current']}])
    claim = next(c for c in ws['claims'] if c['id'] == item['claim_id'])
    assert claim['confirmed'] and claim['refs'] == ['sales_current']


def test_wait_is_idempotent_persistent_and_status_readonly(project):
    store, ws = project
    ws = agent.run_agent(store, ws['id'], ws['revision'])
    saved = (store.folder(ws['id']) / 'state.json').read_bytes()
    restarted = Store(store.root)
    assert agent.run_agent(restarted, ws['id'], ws['revision']) == ws
    assert agent.agent_status(restarted, ws['id'])['phase'] == 'awaiting_decision'
    assert (store.folder(ws['id']) / 'state.json').read_bytes() == saved
    assert approve(restarted, ws)['agent']['phase'] == 'done'


def test_revision_conflicts_invalidate_old_pending(project):
    store, ws = project
    waiting = agent.run_agent(store, ws['id'], ws['revision'])
    changed = store.change(ws['id'], waiting['revision'], {'sales_current': 90})
    with pytest.raises(ValueError, match='其他窗口'):
        approve(store, waiting)
    current = deepcopy(waiting)
    current['revision'] = changed['revision']
    with pytest.raises(ValueError, match='其他窗口'):
        approve(store, current)
    assert agent.agent_status(store, ws['id'])['stale']
    new = agent.run_agent(store, ws['id'], changed['revision'])
    assert new['agent']['run_id'] != waiting['agent']['run_id']


@pytest.mark.parametrize('invalid', ['wrong_kind', 'empty', 'unknown_claim', 'unknown_fact', 'duplicate'])
def test_invalid_decisions_preserve_pending_and_files(project, invalid):
    store, ws = project
    ws = agent.run_agent(store, ws['id'], ws['revision'])
    item = ws['agent']['pending']['items'][0]
    decision = {'kind': 'confirm_links', 'items': [{'claim_id': item['claim_id'], 'refs': item['refs']}]}
    if invalid == 'wrong_kind':
        decision['kind'] = 'approve_repair'
    elif invalid == 'empty':
        decision['items'] = []
    elif invalid == 'unknown_claim':
        decision['items'][0]['claim_id'] = 'invented'
    elif invalid == 'unknown_fact':
        decision['items'][0]['refs'] = ['invented']
    else:
        decision['items'] *= 2
    before = store.read(ws['id'])
    files = file_bytes(store, ws['id'])
    with pytest.raises(ValueError):
        agent.decide(store, ws['id'], ws['revision'], [decision])
    assert store.read(ws['id']) == before
    assert file_bytes(store, ws['id']) == files


def escalate_all(monkeypatch):
    """把全部未确认论断伪装成困难样本，单独验证模型候选的安全处理。

    分级路由下规则唯一项不会触发外部请求（见 test_rule_unique_claims_skip_model），
    因此模型候选的边界必须显式构造困难样本来覆盖。
    """
    routing = agent._candidate_buckets
    def all_hard(ws, rules):
        unique, multi, zero = routing(ws, rules)
        return [], [], unique + multi + zero
    monkeypatch.setattr(agent, '_candidate_buckets', all_hard)


def test_rule_unique_claims_skip_model(project, monkeypatch):
    """规则唯一项不消耗模型调用：演示样例全部为规则唯一，外部请求应为 0 次。"""
    store, ws = project
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'test-secret')
    def unexpected(*args, **kwargs):
        pytest.fail('规则唯一项不应发起模型请求')
    monkeypatch.setattr(httpx, 'post', unexpected)
    ws = agent.run_agent(store, ws['id'], ws['revision'])
    routing = ws['agent']['model_routing']
    assert routing['unique_rule_claims'] == 10 and routing['api_candidate_claims'] == 0
    assert routing['api_suggestions'] == 0
    assert any(t['node'] == 'model_skipped' for t in ws['agent']['trace'])
    assert ws['agent']['pending']['kind'] == 'confirm_links'


def test_unique_unverifiable_source_escalates_but_stale_value_stays_in_rules(project):
    store, public = project
    ws = store.read(public['id'])
    quote = deepcopy(next(c for c in ws['claims'] if c['kind'] == 'quote'))
    ws['claims'] = [quote]
    rules = [{'claim_id': quote['id'], 'refs': quote['refs']}]
    current = next(f for f in ws['facts'] if f['id'] == quote['refs'][0])
    current['value'] = 90
    assert agent._candidate_buckets(ws, rules)[0] == [quote]
    current['unit'] = '件'
    unique, multi, zero = agent._candidate_buckets(ws, rules)
    assert unique == [] and multi == [quote] and zero == []


def test_local_single_source_returns_concrete_rejection(project, monkeypatch):
    store, public = project
    ws = store.read(public['id'])
    quote = deepcopy(next(c for c in ws['claims'] if c['kind'] == 'quote'))
    ws['claims'] = [quote]
    next(f for f in ws['facts'] if f['id'] == quote['refs'][0])['unit'] = '件'
    ws['model_mode'] = 'local'
    monkeypatch.setattr(agent.reranker, 'status', lambda: {'enabled': True})
    result = agent.suggest_links_local(store, ws, {})
    assert not result['suggestions']
    assert result['decisions'][quote['id']]['action'] == 'abstain'
    assert '单位不兼容' in result['decisions'][quote['id']]['reason']


def test_model_trace_and_candidates_are_safe(project, monkeypatch):
    store, ws = project
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'test-secret')
    escalate_all(monkeypatch)
    calls = []
    first = ws['claims'][0]
    def fake_post(url, **kwargs):
        calls.append(url)
        assert kwargs['json']['model'] == 'deepseek-flash'
        # 工具开始轨迹应在外部请求之前落盘。
        trace = store.read(ws['id'])['agent']['trace'][-1]
        assert trace['node'] == 'suggest_links_llm' and trace['model'] == 'deepseek-flash'
        payload = {'suggestions': [{'claim_id': first['id'], 'refs': first['refs'], 'reason': 'test-secret'},
                                   {'claim_id': first['id'], 'refs': ['invented'], 'reason': '坏候选'}]}
        return httpx.Response(200, json={'choices': [{'message': {'content': json.dumps(payload)}}]})
    monkeypatch.setattr(httpx, 'post', fake_post)
    ws = agent.run_agent(store, ws['id'], ws['revision'])
    assert len(calls) == 1
    assert ws['agent']['pending']['kind'] == 'confirm_links'
    trace = next(t for t in ws['agent']['trace'] if t['model'])
    assert trace['model'] == 'deepseek-flash' and '收到 1 项' in trace['summary']
    assert ws['model_call_summary']['remote'] == 1
    assert ws['model_calls'][0]['node'] == 'link_agent'
    assert ws['model_calls'][0]['outcome'] == 'succeeded'
    assert 'test-secret' not in json.dumps(ws)
    assert 'test-secret' not in (store.folder(ws['id']) / 'state.json').read_text(encoding='utf-8')
    approve(store, ws)
    assert len(calls) == 1


def test_model_failure_falls_back_without_secrets(project, monkeypatch):
    store, ws = project
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'test-secret')
    escalate_all(monkeypatch)
    def failed(*args, **kwargs):
        raise httpx.ReadTimeout('provider-body test-secret')
    monkeypatch.setattr(httpx, 'post', failed)
    ws = agent.run_agent(store, ws['id'], ws['revision'])
    assert ws['agent']['pending']['kind'] == 'confirm_links'
    assert any(a['event'] == 'Agent 模型降级' for a in ws['audit'])
    assert 'test-secret' not in json.dumps(ws) and 'provider-body' not in json.dumps(ws)


def test_tool_failure_keeps_prior_approval_and_returns_idle(project, monkeypatch):
    store, ws = project
    ws = agent.run_agent(store, ws['id'], ws['revision'])
    def failed(*args):
        raise RuntimeError('provider-body test-secret')
    monkeypatch.setitem(agent.TOOLS, 'check_claim', failed)
    ws = approve(store, ws)
    assert ws['agent']['phase'] == 'idle'
    assert all(c['confirmed'] for c in ws['claims'])
    assert any(a['event'] == 'Agent 失败' for a in ws['audit'])
    assert 'test-secret' not in json.dumps(ws)


def test_invalid_anchor_blocks_agent(project):
    store, ws = project
    raw = store.read(ws['id'])
    raw['claims'][0]['start'] = -1
    store.write(raw)
    ws = agent.run_agent(store, ws['id'], ws['revision'])
    assert ws['agent']['phase'] == 'idle'
    assert not any(c['confirmed'] for c in ws['claims'])


def test_tool_error_after_commit_does_not_overwrite_decision(project, monkeypatch):
    store, ws = project
    ws = agent.run_agent(store, ws['id'], ws['revision'])
    def commit_then_fail(store, current, payload):
        agent.confirm_links(store, current, payload)
        raise RuntimeError('模拟持久化之后的错误')
    monkeypatch.setitem(agent.TOOLS, 'confirm_links', commit_then_fail)
    ws = approve(store, ws)
    assert ws['agent']['phase'] == 'idle'
    assert all(c['confirmed'] for c in store.read(ws['id'])['claims'])


def test_model_disagreement_pauses_for_choice(project, monkeypatch):
    store, ws = project
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'test-secret')
    escalate_all(monkeypatch)
    claim = next(c for c in ws['claims'] if c['kind'] == 'quote')
    payload = {'suggestions': [{'claim_id': claim['id'], 'refs': ['sales_prev'], 'reason': '不同候选'}]}
    monkeypatch.setattr(httpx, 'post', lambda *args, **kwargs: httpx.Response(
        200, json={'choices': [{'message': {'content': json.dumps(payload)}}]}))
    ws = agent.run_agent(store, ws['id'], ws['revision'])
    assert ws['agent']['pending']['kind'] == 'resolve_ambiguity'
    item = next(i for i in ws['agent']['pending']['items'] if i['claim_id'] == claim['id'])
    assert len(item['options']) == 2
    assert next(c for c in ws['claims'] if c['id'] == claim['id'])['refs'] == claim['refs']


def test_agent_api_flow_and_boundaries(tmp_path, monkeypatch):
    client = TestClient(app_module.create_app(tmp_path / 'api'))
    ws = client.post('/api/projects/demo').json()
    base = f'/api/projects/{ws["id"]}/agent'
    assert client.get(base + '/status').json()['phase'] == 'idle'
    assert client.post(base + '/run', json={'revision': ws['revision']}, headers={'Origin': 'https://evil.example'}).status_code == 403
    response = client.post(base + '/run', json={'revision': ws['revision']})
    assert response.status_code == 200 and response.headers['cache-control'] == 'no-store'
    waiting = response.json()
    pending = waiting['agent']['pending']
    payload = {'revision': waiting['revision'], 'decisions': [{'kind': pending['kind'], 'items': [
        {'claim_id': i['claim_id'], 'refs': i['refs']} for i in pending['items']]}]}
    assert client.post(base + '/decide', json=payload).json()['agent']['phase'] == 'done'
    assert client.post(base + '/decide', json=payload).status_code == 400
    monkeypatch.setenv('ZHILIAN_ACCESS_PASSWORD', 'agent-password')
    for endpoint in ('run', 'decide'):
        assert client.post(base + '/' + endpoint, json=payload).status_code == 401
    assert client.get(base + '/status').status_code == 401
    assert client.get(base + '/status', auth=('zhilian', 'agent-password')).status_code == 200


def test_uploaded_documents_can_run_agent_without_demo_mode(tmp_path):
    paths = create_demo(tmp_path / 'input')
    wb = load_workbook(paths[0])
    wb.active['E3'] = 90
    wb.save(paths[0])
    wb.close()
    client = TestClient(app_module.create_app(tmp_path / 'api'))
    response = client.post('/api/projects', data={'name': '自有文件智能体测试'}, files=[
        ('files', (p.name, p.read_bytes(), 'application/octet-stream')) for p in paths])
    assert response.status_code == 200
    ws = response.json()
    assert not ws['demo'] and ws['summary']['inconsistent'] > 0
    base = f'/api/projects/{ws["id"]}/agent'
    ws = client.post(base + '/run', json={'revision': ws['revision']}).json()
    assert ws['agent']['pending']['kind'] == 'confirm_links'
    for expected in ('confirm_links', 'approve_repair'):
        pending = ws['agent']['pending']
        assert pending['kind'] == expected
        response = client.post(base + '/decide', json={'revision': ws['revision'], 'decisions': [{
            'kind': expected, 'items': [{'claim_id': i['claim_id'], 'refs': i['refs']} for i in pending['items']]}]})
        assert response.status_code == 200
        ws = response.json()
    assert ws['agent']['phase'] == 'done' and ws['agent']['result']['repaired'] > 0
    assert ws['summary']['inconsistent'] == 0


def test_agent_exports_all_current_office_files(tmp_path):
    client = TestClient(app_module.create_app(tmp_path / 'api'))
    ws = client.post('/api/projects/demo').json()
    base = f'/api/projects/{ws["id"]}'
    originals = {d['kind']: client.get(d['download_url']).content for d in ws['documents']}
    ws = client.post(base + '/facts', json={'revision': ws['revision'],
        'values': {'sales_current': 90, 'product_a': 60, 'spending': 110}}).json()
    ws = client.post(base + '/agent/run', json={'revision': ws['revision']}).json()
    for kind in ('confirm_links', 'approve_repair'):
        pending = ws['agent']['pending']
        assert pending['kind'] == kind
        response = client.post(base + '/agent/decide', json={'revision': ws['revision'], 'decisions': [{
            'kind': kind, 'items': [{'claim_id': i['claim_id'], 'refs': i['refs']} for i in pending['items']]}]})
        assert response.status_code == 200
        ws = response.json()
    assert ws['agent']['phase'] == 'done' and ws['agent']['result']['repaired'] == 11

    response = client.get(base + '/export')
    assert response.status_code == 200
    downloads = {}
    raw = client.app.state.store.read(ws['id'])
    generation = client.app.state.store.folder(ws['id']) / raw['generation']
    with ZipFile(BytesIO(response.content)) as archive:
        assert set(archive.namelist()) == {d['name'] for d in ws['documents']} | {'核验报告.md', '知链核验记录.json'}
        assert json.loads(archive.read('知链核验记录.json'))['revision'] == ws['revision']
        for d in ws['documents']:
            response = client.get(d['download_url'])
            assert response.status_code == 200 and response.content
            assert response.headers['cache-control'] == 'no-store'
            stored = next(item['stored_name'] for item in raw['documents'] if item['id'] == d['id'])
            assert response.content == archive.read(d['name']) == (generation / stored).read_bytes()
            assert response.content != originals[d['kind']]
            downloads[d['kind']] = response.content
    assert set(downloads) == {'xlsx', 'docx', 'pptx'}
    workbook = load_workbook(BytesIO(downloads['xlsx']))
    assert [workbook.active[cell].value for cell in ('E3', 'E4', 'E6')] == [90, 60, 110]
    workbook.close()
    doc = Document(BytesIO(downloads['docx']))
    assert '本期销售额为90万元' in '\n'.join(p.text for p in doc.paragraphs)
    deck = Presentation(BytesIO(downloads['pptx']))
    text = '\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    assert '本期销售额为90万元' in text and '较上期下降10%' in text
    chart = next(s.chart for slide in deck.slides for s in slide.shapes if s.has_chart)
    assert list(chart.series[0].values) == [100, 90]
