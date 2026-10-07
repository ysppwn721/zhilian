import json
from pathlib import Path

import pytest

fitz = pytest.importorskip("fitz")

from zhilian import pdf_office_export


def _make_pdf(path):
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Revenue for current period is 125 CNY.", fontsize=12)
    document.save(path)
    document.close()


def test_export_pdf_to_office_creates_provenance_files(tmp_path, monkeypatch):
    source = tmp_path / "annual_report.pdf"
    _make_pdf(source)
    monkeypatch.setattr(pdf_office_export, "_pdf_text_pages", lambda path: [
        {"page": 1, "text": "Revenue for current period is 125 CNY."}
    ])
    monkeypatch.setattr(pdf_office_export, "extract_pdf", lambda *args, **kwargs: {
        "pages": 1, "text_layer_pages": 1, "text_layer_ratio": 1.0,
        "facts": [{"id": "pdf-p1-t1-r2-current", "metric": "营业收入", "period": "本期",
                   "value": 125.0, "unit": "元", "scope": "annual_report.pdf#p1",
                   "source_file": "annual_report.pdf", "source_page": 1,
                   "source_table": 1, "source_row": 2}],
        "tables": [{"page": 1, "table": 1, "rows": [["指标", "本期", "上期"]],
                    "validation": {"checks": []}}],
        "chart_candidates": [],
    })

    manifest = pdf_office_export.export_pdf_to_office(source, tmp_path / "out")

    word = tmp_path / "out" / "annual_report_正文提取.docx"
    excel = tmp_path / "out" / "annual_report_事实表候选.xlsx"
    assert word.exists() and excel.exists()
    assert manifest["original_pdf_modified"] is False
    assert manifest["requires_confirmation"] is True
    assert manifest["text_anchor_count"] == 1
    assert manifest["claim_candidate_count"] == 1
    assert (tmp_path / "out" / "annual_report_论断候选.jsonl").read_text(encoding="utf-8").strip()
    anchors = json.loads((tmp_path / "out" / "annual_report_文字锚点.json").read_text(encoding="utf-8"))
    assert anchors["anchors"][0]["page"] == 1
    assert "Revenue for current period" in anchors["anchors"][0]["text"]
    assert json.loads((tmp_path / "out" / "annual_report_转换清单.json").read_text(encoding="utf-8"))["facts"] == 1

    from openpyxl import load_workbook
    workbook = load_workbook(excel, read_only=True, data_only=True)
    assert workbook["事实候选"]["A2"].value == "pdf-p1-t1-r2-current"
    assert workbook["事实候选"]["H2"].value == "annual_report.pdf"
    assert workbook["事实候选"]["I2"].value == 1
    assert workbook["事实候选"]["L2"].value == "待复核：缺少独立交叉校验"
    assert workbook["冲突待复核"].max_row == 1
    workbook.close()


def test_pdf_raw_table_sheet_does_not_trigger_fact_table_row_limit(tmp_path):
    from openpyxl import Workbook
    from zhilian.office import HEADERS, read_facts

    path = tmp_path / "pdf_candidates.xlsx"
    workbook = Workbook()
    facts = workbook.active
    facts.title = "事实候选"
    facts.append(HEADERS)
    facts.append(["f1", "皖能电力", "营业收入", "本期", 100, "元", "待确认"])
    raw = workbook.create_sheet("表格原文")
    raw.append(["来源文件", "来源页", "表格编号", "行号", "列号", "原始内容"])
    for row in range(2003):
        raw.append(["report.pdf", 1, 1, row + 1, 1, "raw cell"])
    workbook.save(path)
    workbook.close()

    result = read_facts(path, "source")

    assert len(result) == 1
    assert result[0]["metric"] == "营业收入"


def test_pdf_project_endpoint_routes_to_office_conversion(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from zhilian.app import create_app

    source = tmp_path / "input.pdf"
    _make_pdf(source)
    app = create_app(tmp_path / "api")
    called = {}

    def fake_create(name, path, subject=None):
        called.update(name=name, suffix=path.suffix, subject=subject, exists=path.exists())
        return {"id": "a" * 32, "name": name}

    monkeypatch.setattr(app.state.store, "create_from_pdf", fake_create)
    response = TestClient(app).post(
        "/api/projects/from-pdf",
        data={"name": "PDF 项目", "subject": "某公司"},
        files={"file": (source.name, source.read_bytes(), "application/pdf")},
    )
    assert response.status_code == 200, response.text
    assert called == {"name": "PDF 项目", "suffix": ".pdf", "subject": "某公司", "exists": True}


def test_store_keeps_original_pdf_and_archives_it(tmp_path, monkeypatch):
    from zipfile import ZipFile
    from zhilian.demo import create_demo
    from zhilian.store import Store
    import zhilian.store as store_module

    source = tmp_path / "source.pdf"
    source.write_bytes(b"pdf-origin")

    def fake_export(pdf_path, output_dir, subject=None):
        paths = create_demo(output_dir)
        excel = next(path for path in paths if path.suffix == ".xlsx")
        word = next(path for path in paths if path.suffix == ".docx")
        return {"source_pdf": str(pdf_path), "word": str(word), "excel": str(excel),
                "facts": 6, "tables": 1, "chart_candidates": 0, "text_pages": 1,
                "requires_confirmation": True, "original_pdf_modified": False}

    monkeypatch.setattr(store_module, "export_pdf_to_office", fake_export)
    store = Store(tmp_path / "data")
    result = store.create_from_pdf("PDF 项目", source, subject="主体")
    state = store.read(result["id"])
    assert len(state["pdf_origins"]) == 1
    archive = store.archive(result["id"])
    with ZipFile(archive) as zipped:
        assert "原始PDF/source.pdf" in zipped.namelist()


def test_store_imports_searchable_pdf_when_no_fact_rows_are_admitted(tmp_path, monkeypatch):
    from docx import Document
    from openpyxl import Workbook
    from zhilian.office import HEADERS
    from zhilian.store import Store
    import zhilian.store as store_module

    source = tmp_path / "text_only.pdf"
    source.write_bytes(b"pdf-origin")

    def fake_export(pdf_path, output_dir, subject=None):
        output_dir = Path(output_dir)
        excel = output_dir / "candidate.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "事实候选"
        sheet.append(HEADERS)
        workbook.save(excel)
        word = output_dir / "body.docx"
        document = Document()
        document.add_paragraph("可提取的 PDF 正文段落。")
        document.save(word)
        return {"source_pdf": str(pdf_path), "word": str(word), "excel": str(excel),
                "facts": 0, "tables": 1, "chart_candidates": 0, "text_pages": 1,
                "text_layer_pages": 1, "requires_confirmation": True,
                "original_pdf_modified": False}

    monkeypatch.setattr(store_module, "export_pdf_to_office", fake_export)
    store = Store(tmp_path / "data")
    result = store.create_from_pdf("PDF 项目", source)
    state = store.read(result["id"])

    assert state["facts"] == []
    assert state["allow_empty_facts"] is True
    assert state["fact_quality"]["status"] == "blocked"
    assert len(state["pdf_origins"]) == 1


def test_store_exports_independent_revised_pdf_and_manifest(tmp_path, monkeypatch):
    from zhilian.demo import create_demo
    from zhilian.store import Store
    import zhilian.store as store_module

    source = tmp_path / "source.pdf"
    source.write_bytes(b"pdf-origin")

    def fake_export(pdf_path, output_dir, subject=None):
        paths = create_demo(output_dir)
        excel = next(path for path in paths if path.suffix == ".xlsx")
        word = next(path for path in paths if path.suffix == ".docx")
        return {"source_pdf": str(pdf_path), "word": str(word), "excel": str(excel),
                "facts": 6, "tables": 1, "chart_candidates": 0,
                "requires_confirmation": True, "original_pdf_modified": False}

    def fake_docx_to_pdf(docx, output):
        output.write_bytes(b"revised-pdf")
        return {"docx": str(docx), "pdf": str(output), "backend": "test",
                "docx_sha256": "docx", "pdf_sha256": "pdf", "original_is_read_only": True}

    def fake_verify(original, revised):
        return {"passed": True, "status": "passed", "page_count": {"unchanged": True}}

    def fake_manifest(original, exported, verification, output):
        output.write_text(json.dumps(verification), encoding="utf-8")
        return verification

    monkeypatch.setattr(store_module, "export_pdf_to_office", fake_export)
    monkeypatch.setattr(store_module, "export_docx_to_pdf", fake_docx_to_pdf)
    monkeypatch.setattr(store_module, "verify_revised_pdf", fake_verify)
    monkeypatch.setattr(store_module, "compare_pdf_visual_fidelity", lambda original, revised, **kwargs: {
        "passed": True, "status": "passed", "changed_page_count": 0,
        "page_count": {"original": 1, "revised": 1, "unchanged": True},
    })
    monkeypatch.setattr(store_module, "build_revised_pdf_manifest", fake_manifest)
    store = Store(tmp_path / "data")
    result = store.create_from_pdf("PDF 项目", source, subject="主体")
    state = store.read(result["id"])
    revised = store.export_revised_pdf(result["id"], state["revision"])
    saved = store.read(result["id"])
    assert revised["pdf_revision"]["verification"]["status"] == "passed"
    assert revised["pdf_revision"]["verification"]["visual_check"]["passed"] is True
    assert len(saved["pdf_revisions"]) == 1
    assert (store.folder(result["id"]) / saved["generation"] / saved["pdf_revisions"][0]["stored_name"]).read_bytes() == b"revised-pdf"
